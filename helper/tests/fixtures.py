"""Synthetic DOCX / PPTX files covering the text locations phase 1 handles. No third-party content."""
import docx
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WPS_URI = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
# "This is a pathology diagnosis." in Chinese, and a clone name with a Greek letter (notation, not language)
CHINESE = "".join(chr(c) for c in (0x8FD9, 0x662F, 0x75C5, 0x7406, 0x8BCA, 0x65AD, 0x3002))
GREEK_CLONE = "34" + chr(0x3B2) + "E12"
MICRO_UNIT = "5 " + chr(0xB5) + "m"  # MICRO SIGN, not the Greek mu
PRIME_END = "3" + chr(0x2B9)  # MODIFIER LETTER PRIME (category Lm), as in 3' end
SHORT_PARAGRAPHS = ("52", "+", "<0.001", "No", "Yes", CHINESE, GREEK_CLONE, MICRO_UNIT, PRIME_END)

_TEXT_BOX = (
    '<w:r xmlns:w="' + W_NS + '" xmlns:a="' + A_NS + '" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:wps="' + WPS_URI + '"><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
    '<wp:extent cx="2743200" cy="457200"/><wp:docPr id="100" name="Text Box 1"/>'
    '<a:graphic><a:graphicData uri="' + WPS_URI + '"><wps:wsp><wps:cNvSpPr txBox="1"/>'
    '<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2743200" cy="457200"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>'
    '<wps:txbx><w:txbxContent><w:p><w:r><w:t/></w:r></w:p></w:txbxContent></wps:txbx>'
    '<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')


def add_text_box_paragraph(d, text):
    """Body paragraph holding only an inline text box whose one paragraph reads `text`."""
    p = d.add_paragraph()
    run = etree.fromstring(_TEXT_BOX)
    run.find(".//{%s}t" % W_NS).text = text
    p._p.append(run)
    return p


def _send_to_back(shape):
    el = shape._element
    tree = el.getparent()
    tree.remove(el)
    tree.insert(2, el)  # after p:nvGrpSpPr and p:grpSpPr


def reorder_slides(path, order):
    """Rewrite p:sldIdLst so the slide at old position order[k] shows as slide k + 1 (part names stay)."""
    prs = Presentation(path)
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for el in ids:
        lst.remove(el)
    for n in order:
        lst.append(ids[n])
    prs.save(path)
    return path


def make_pptx(path):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Meningioma grading"
    body = s1.placeholders[1].text_frame
    body.text = "Grade depends on mitoses"
    p = body.add_paragraph()
    p.add_run().text = "In "
    bold = p.add_run()
    bold.text = "NF2-altered"
    bold.font.bold = True
    p.add_run().text = " meningiomas, 1p loss matters"
    p.add_line_break()
    link = p.add_run()
    link.text = "See the WHO site"
    link.hyperlink.address = "https://www.who.int"
    fld = etree.Element("{%s}fld" % A_NS, id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
    etree.SubElement(fld, "{%s}rPr" % A_NS, lang="en-US")
    etree.SubElement(fld, "{%s}t" % A_NS).text = "1"
    end = p._p.find("{%s}endParaRPr" % A_NS)
    if end is not None:
        end.addprevious(fld)
    else:
        p._p.append(fld)
    table = s1.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "Marker"
    table.cell(0, 1).text = "Result"
    table.cell(1, 0).text = "BRAF VE1"
    table.cell(1, 1).text = "Negative staining"
    group = s1.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(6), Inches(3), Inches(0.5))
    box.text_frame.text = "Atypical features are present"
    s1.notes_slide.notes_text_frame.text = "Speaker note about grading"
    # references slide: the title placeholder is a stop heading, though a text box comes first in shape order
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "References"
    s2.placeholders[1].text_frame.text = "Louis DN et al. Neuro Oncol 2021"
    banner = s2.shapes.add_textbox(Inches(1), Inches(7), Inches(4), Inches(0.4))
    banner.text_frame.text = "15th Pathology Conference"
    _send_to_back(banner)
    # titled content slide: a "References:" footnote sent to the back (first in shape order) and a body placeholder
    # line "References" do not stop it; its notes start with "References"
    s3 = prs.slides.add_slide(prs.slide_layouts[1])
    s3.shapes.title.text = "Survival by grade"
    content = s3.placeholders[1].text_frame
    content.text = "Grade 3 tumours recur early"
    content.add_paragraph().text = "References"
    foot = s3.shapes.add_textbox(Inches(1), Inches(6.5), Inches(6), Inches(0.5))
    foot.text_frame.text = "References:"
    foot.text_frame.add_paragraph().text = "Louis DN et al. Neuro Oncol 2021"
    _send_to_back(foot)
    notes = s3.notes_slide.notes_text_frame
    notes.text = "References"
    notes.add_paragraph().text = "Mention the 2021 criteria"
    # blank slide whose first text is a table column header "Reference"
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    ranges = s4.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(1)).table
    ranges.cell(0, 0).text = "Reference"
    ranges.cell(0, 1).text = "Range"
    ranges.cell(1, 0).text = "Ki-67 index"
    ranges.cell(1, 1).text = "Below 4 percent"
    # title-only slide with the title left empty: a slide number, then "References" as the first line of text
    s5 = prs.slides.add_slide(prs.slide_layouts[5])
    for n, text in enumerate(("52", "References", "Smith J et al. Lancet 2020")):
        s5.shapes.add_textbox(Inches(1), Inches(1 + n), Inches(6), Inches(0.5)).text_frame.text = text
    prs.save(path)
    return path


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def _run(text=None, rstyle=None):
    r = _el("w:r")
    if rstyle:
        rpr = _el("w:rPr")
        rpr.append(_el("w:rStyle", **{"w:val": rstyle}))
        r.append(rpr)
    if text is not None:
        t = _el("w:t")
        t.text = text
        r.append(t)
    return r


def set_outline_level(p, level):
    """Give a python-docx paragraph a direct w:outlineLvl (0-8 = heading levels, 9 = body text)."""
    p._p.get_or_add_pPr().append(_el("w:outlineLvl", **{"w:val": str(level)}))
    return p


def make_docx(path):
    d = docx.Document()
    d.add_heading("Meningioma grading", level=1)
    p = d.add_paragraph()
    p.add_run("In ")
    p.add_run("NF2-altered").bold = True
    p.add_run(" meningiomas")
    p.add_run("1").font.superscript = True
    r = p.add_run()
    r.add_tab()
    r.add_text("criteria")

    p2 = d.add_paragraph("See ")
    rid = d.part.relate_to("https://www.who.int", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = _el("w:hyperlink", **{"r:id": rid})
    link.append(_run("the WHO site", rstyle="Hyperlink"))
    p2._p.insert(1, _el("w:bookmarkStart", **{"w:id": "0", "w:name": "_GoBack"}))
    p2._p.append(_el("w:proofErr", **{"w:type": "spellStart"}))
    p2._p.append(link)
    p2._p.append(_el("w:proofErr", **{"w:type": "spellEnd"}))
    p2.add_run(" for page ")
    for kind in ("begin", None, "separate", "result", "end"):
        fr = _el("w:r")
        if kind is None:
            it = _el("w:instrText")
            it.text = " PAGE "
            fr.append(it)
        elif kind == "result":
            t = _el("w:t")
            t.text = "4"
            fr.append(t)
        else:
            fr.append(_el("w:fldChar", **{"w:fldCharType": kind}))
        p2._p.append(fr)
    p2._p.append(_el("w:bookmarkEnd", **{"w:id": "0"}))
    p2.add_run().add_break()
    p2.add_run("Next line of text")

    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Marker"
    table.cell(0, 1).text = "Result"
    table.cell(1, 0).text = "BRAF VE1"
    table.cell(1, 1).text = "Negative staining"

    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Department of Pathology"
    sec.footer.paragraphs[0].text = "Confidential draft"

    for text in SHORT_PARAGRAPHS:
        d.add_paragraph(text)
    add_text_box_paragraph(d, "Text box caption")

    tracked = d.add_paragraph()
    ins = _el("w:ins", **{"w:id": "9", "w:author": "Reviewer"})
    ins.append(_run("Inserted sentence"))
    tracked._p.append(ins)

    d.add_heading("References", level=1)
    d.add_paragraph("Louis DN et al. Neuro Oncol 2021")
    d.save(path)
    return path
