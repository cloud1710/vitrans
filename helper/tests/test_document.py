import unicodedata
from pathlib import Path

import docx
import pptx
import pytest

from describe import describe
from fixtures import (CHINESE, GREEK_CLONE, MICRO_UNIT, PRIME_END, add_text_box_paragraph, make_docx, make_pptx,
                      reorder_slides, set_outline_level)
from vitranslate.config import REPO_ROOT
from vitranslate.ooxml.codec import apply
from vitranslate.ooxml.document import collect_items, is_back_matter_heading, is_stop_heading
from vitranslate.ooxml.ns import q
from vitranslate.ooxml.package import OoxmlPackage


def _texts(items):
    return [i.segment.plain for i in items if i.segment]


def _by_text(items):
    return {i.text: i for i in items}


def _docx_items(tmp_path, build):
    d = docx.Document()
    build(d)
    path = tmp_path / "t.docx"
    d.save(path)
    return collect_items(OoxmlPackage(path))


def test_stop_heading():
    assert is_stop_heading("References:")
    assert is_stop_heading("REFERENCE")
    assert is_stop_heading("  BIBLIOGRAPHY ")
    assert not is_stop_heading("References for grading")
    assert is_stop_heading("19. BIBLIOGRAPHY")
    assert is_stop_heading("3) References")
    assert is_stop_heading("Selected" + chr(0xA0) + "references")
    assert not is_stop_heading("2021 references")
    for phrase in ("Reference list", "Further reading", "Suggested reading", "Key references", "Works cited"):
        assert is_stop_heading(phrase), phrase
    nfc = "t" + chr(0xE0) + "i li" + chr(0x1EC7) + "u tham kh" + chr(0x1EA3) + "o"
    nfd = unicodedata.normalize("NFD", nfc)
    assert nfd != nfc
    assert is_stop_heading(nfd)
    assert is_stop_heading("5. " + nfd.upper() + ":")


def test_pptx_items(tmp_path):
    items = collect_items(OoxmlPackage(make_pptx(tmp_path / "a.pptx")))
    texts = _texts(items)
    for want in ("Meningioma grading", "Negative staining", "Atypical features are present",
                 "Speaker note about grading"):
        assert want in texts, want
    body = next(i for i in items if i.segment and "NF2-altered" in i.segment.plain)
    assert body.segment.tagged == (
        "In <g1>NF2-altered</g1> meningiomas, 1p loss matters<x1/><g2>See the WHO site</g2><x2/>")
    assert body.location == "Slide 1"
    refs = [i for i in items if i.part.endswith("slides/slide2.xml")]
    assert refs and all(i.segment is None for i in refs)
    assert any(i.skip == "references" for i in refs)
    # layout / master placeholder prompts never become segments
    assert not any("Click to edit" in t for t in texts)


def test_pptx_references_stop(tmp_path):
    items = collect_items(OoxmlPackage(make_pptx(tmp_path / "a.pptx")))

    def skips(loc):
        return {i.skip for i in items if i.location == loc}

    # title placeholder "References" (a text box comes first in shape order): whole slide
    assert skips("Slide 2") == {"references"}
    # empty title, slide number "52", then a "References" text box as the first line of text: whole slide
    assert skips("Slide 5") == {"references"}
    # titled slide: a "References:" footnote first in shape order and a body placeholder line "References" do not
    # stop it; nor do "References" as the first line of notes, or a "Reference" table header on an untitled slide
    for loc, wants in (("Slide 3", ("Survival by grade", "References:", "References")),
                       ("Notes 3", ("References", "Mention the 2021 criteria")),
                       ("Slide 4", ("Reference", "Ki-67 index"))):
        assert "references" not in skips(loc), loc
        plains = [i.segment.plain for i in items if i.location == loc and i.segment]
        for want in wants:
            assert want in plains, (loc, want)


def test_locations_follow_presentation_order(tmp_path):
    path = reorder_slides(make_pptx(tmp_path / "a.pptx"), [4, 3, 2, 1, 0])
    items = collect_items(OoxmlPackage(path))
    by = _by_text(items)
    assert by["Meningioma grading"].location == "Slide 5"
    assert by["Speaker note about grading"].location == "Notes 5"
    assert by["Survival by grade"].location == "Slide 3"
    assert items[0].part == "ppt/slides/slide5.xml" and items[0].location == "Slide 1"


def test_docx_items(tmp_path):
    items = collect_items(OoxmlPackage(make_docx(tmp_path / "a.docx")))
    texts = _texts(items)
    for want in ("Meningioma grading", "Department of Pathology", "Confidential draft", "Negative staining"):
        assert want in texts, want
    para = next(i for i in items if i.segment and "NF2-altered" in i.segment.plain)
    assert para.segment.tagged == "In <g1>NF2-altered</g1> meningiomas<g2>1</g2><x1/>criteria"
    link = next(i for i in items if i.segment and "WHO site" in i.segment.plain)
    assert link.segment.tagged == "See <g1>the WHO site</g1> for page <x1/><x2/>Next line of text"
    assert any(i.skip == "tracked changes" for i in items)
    # references: the heading and everything after it in the body, nothing before it
    body = [i for i in items if i.part == "word/document.xml"]
    heading = next(n for n, i in enumerate(body) if i.text == "References")
    assert not any(i.skip == "references" for i in body[:heading])
    assert [i.skip for i in body[heading:]] == ["references", "references"]
    assert body[-1].text.startswith("Louis DN")
    by = _by_text(items)
    for text in ("52", "+", "<0.001", GREEK_CLONE, MICRO_UNIT, PRIME_END):
        assert by[text].skip == "no_text" and by[text].segment is None, text
    for text in ("No", "Yes"):
        assert by[text].segment and by[text].segment.plain == text, text
    assert by[CHINESE].skip == "non_latin" and by[CHINESE].segment is None
    # a paragraph holding only a text box has no text of its own; the text box paragraph is translated
    box = by["Text box caption"]
    assert box.segment and box.segment.plain == "Text box caption"
    outer = next(i for i in body if i.element is not box.element
                 and any(e is box.element for e in i.element.iter(q("w:p"))))
    assert outer.skip == "no_text" and outer.segment is None


def test_docx_reference_column_does_not_stop(tmp_path):
    def build(d):
        lab = d.add_table(rows=2, cols=3)
        for c, text in enumerate(("Test", "Result", "Reference")):
            lab.cell(0, c).text = text
        for c, text in enumerate(("Ki-67 index", "12 percent", "Below 4 percent")):
            lab.cell(1, c).text = text
        d.add_paragraph("Final diagnosis: atypical meningioma.")

    items = _docx_items(tmp_path, build)
    assert not any(i.skip == "references" for i in items)
    assert _by_text(items)["Final diagnosis: atypical meningioma."].segment


def test_docx_text_box_does_not_stop(tmp_path):
    def build(d):
        add_text_box_paragraph(d, "References")
        d.add_paragraph("Final diagnosis: atypical meningioma.")

    items = _docx_items(tmp_path, build)
    assert not any(i.skip == "references" for i in items)
    assert _by_text(items)["Final diagnosis: atypical meningioma."].segment


@pytest.mark.parametrize("style, outline", [
    ("Heading 1", None),   # heading-styled stop: the region ends at the next heading
    ("Title", None),
    (None, 0),             # direct outline level 0 = heading
    (None, None),          # plain paragraph "References": a back-matter title ends it (his report 2026-10-15)
    (None, 9),             # outline level 9 = body text: likewise
])
def test_docx_references_region(tmp_path, style, outline):
    def build(d):
        d.add_paragraph("Body text before the list")
        head = d.add_paragraph("References", style=style)
        if outline is not None:
            set_outline_level(head, outline)
        d.add_paragraph("Louis DN et al. Neuro Oncol 2021")
        d.add_heading("Appendix A", level=1)
        d.add_paragraph("Staining protocol for the lab")

    by = _by_text(_docx_items(tmp_path, build))
    assert by["Body text before the list"].segment
    assert by["References"].skip == "references"
    assert by["Louis DN et al. Neuro Oncol 2021"].skip == "references"
    for text in ("Appendix A", "Staining protocol for the lab"):
        assert by[text].segment and not by[text].skip, text


BACK_MATTER = ["Appendix", "Appendix A", "APPENDIX 2", "Appendix S1: Staining protocols", "Appendices", "Annex 1",
               "Supplementary", "Supplementary Material", "Supplementary data", "Supplementary Information",
               "Supplementary methods", "Supplementary Figures", "Supplementary tables", "Supporting Information",
               "Online Resource 1", "Acknowledgements", "Acknowledgments", "Abbreviations", "Glossary",
               "Figure legends", "Tables", "Author contributions", "Funding", "Conflict of interest",
               "Conflicts of Interest", "Competing interests", "Data availability", "About the author",
               "About the Authors", "7. Acknowledgements", "A. Appendix", "Appendix 1. Questionnaire", "Funding:",
               "Phụ lục", "PHỤ LỤC 2", "Phu luc 1"]


def test_back_matter_heading():
    for text in BACK_MATTER:
        assert is_back_matter_heading(text), text
    for text in ["References", "Staining protocol for the lab", "Louis DN et al. Neuro Oncol 2021",
                 "Funding. J Synth Path. 2020;1:1-9.", "Tables of the synthetic series in a long form for the lab",
                 "Appendix A. Example AB et al. 2019", "Appendicitis", "Funding was provided by the synthetic fund",
                 "Abbreviations used in the synthetic text are listed below in full"]:
        assert not is_back_matter_heading(text), text


@pytest.mark.parametrize("title", BACK_MATTER)
@pytest.mark.parametrize("stop", [("style", "Normal"), ("outline", 9), ("style", "Heading 1"), ("outline", 0)])
def test_docx_back_matter_closes_references(tmp_path, title, stop):
    """A back-matter title closes the references region whatever its style or level: a plain paragraph after a
    plain "References", a deeper heading after a level-0 one. The title and what follows are translated, the
    entries before it are not, nor a citation-looking paragraph."""
    def build(d):
        _para(d, "References", stop)
        d.add_paragraph("Louis DN et al. Neuro Oncol 2021")
        d.add_paragraph("Funding. J Synth Path. 2020;1:1-9.")
        _para(d, title, ("style", "Normal") if stop[0] == "style" and stop[1] == "Normal" else H3)
        d.add_paragraph("Staining protocol for the lab")

    by = _by_text(_docx_items(tmp_path, build))
    for text in ("References", "Louis DN et al. Neuro Oncol 2021", "Funding. J Synth Path. 2020;1:1-9."):
        assert by[text].skip == "references", text
    for text in (title, "Staining protocol for the lab"):
        assert by[text].segment and not by[text].skip, text


def test_docx_plain_paragraph_after_plain_references_stays_inside(tmp_path):
    def build(d):
        d.add_paragraph("References")
        d.add_paragraph("Louis DN et al. Neuro Oncol 2021")
        d.add_paragraph("Perry A, Example B. A synthetic study of the lab. Brain Pathol 2020;30:1-9.")
        d.add_paragraph("Staining protocol for the lab")

    items = _docx_items(tmp_path, build)
    assert all(i.skip == "references" for i in items)


def _para(d, text, how):
    """how: ("style", python-docx style name), ("outline", direct w:outlineLvl), ("style+outline", (name, level))
    or ("id", a pStyle id that styles.xml does not define)."""
    kind, value = how
    if kind == "style":
        return d.add_paragraph(text, style=value)
    if kind == "style+outline":
        return set_outline_level(d.add_paragraph(text, style=value[0]), value[1])
    p = d.add_paragraph(text)
    if kind == "outline":
        set_outline_level(p, value)
    else:
        p._p.style = value
    return p


H1, H2, H3 = ("style", "Heading 1"), ("style", "Heading 2"), ("style", "Heading 3")


@pytest.mark.parametrize("stop, sub, appendix", [
    (H1, H2, H1),
    (("outline", 0), ("outline", 1), ("outline", 0)),
    (("style", "Title"), H2, H1),               # Title and Heading 1 are both level 0
    (("outline", 0), H2, H1),
    (H1, H2, ("id", "AppendixHeading")),        # heading-like style without a level: the stop heading's level
    (("style+outline", ("Heading 2", 0)), H2, H1),  # a direct outline level wins over the style
])
def test_docx_references_subheadings(tmp_path, stop, sub, appendix):
    def build(d):
        _para(d, "References", stop)
        _para(d, "Primary sources", sub)
        d.add_paragraph("Louis DN et al. Neuro Oncol 2021")
        _para(d, "Reviews", sub)
        d.add_paragraph("Perry A et al. Brain Pathol 2020")
        _para(d, "Appendix A", appendix)
        d.add_paragraph("Staining protocol for the lab")

    by = _by_text(_docx_items(tmp_path, build))
    for text in ("References", "Primary sources", "Louis DN et al. Neuro Oncol 2021", "Reviews",
                 "Perry A et al. Brain Pathol 2020"):
        assert by[text].skip == "references", text
    for text in ("Appendix A", "Staining protocol for the lab"):
        assert by[text].segment and not by[text].skip, text


def test_docx_stop_heading_without_level_closes_at_any_heading(tmp_path):
    def build(d):
        _para(d, "References", ("id", "ReferencesHeading"))
        d.add_paragraph("Louis DN et al. Neuro Oncol 2021")
        _para(d, "Appendix A", H3)
        d.add_paragraph("Staining protocol for the lab")

    by = _by_text(_docx_items(tmp_path, build))
    assert by["Louis DN et al. Neuro Oncol 2021"].skip == "references"
    assert by["Appendix A"].segment and by["Staining protocol for the lab"].segment


def _identity_roundtrip(path, tmp_path):
    pkg = OoxmlPackage(path)
    items = [i for i in collect_items(pkg) if i.segment]
    before = [describe(i.element) for i in items]
    for i in items:
        apply(i.element, i.segment, i.segment.tagged)
        pkg.mark_dirty(i.part)
    after = [describe(i.element) for i in items]
    for b, a, i in zip(before, after, items):
        assert a == b, (i.location, i.segment.tagged)
    out = tmp_path / ("out" + Path(path).suffix)
    pkg.save(out)
    return out


def test_identity_roundtrip_fixtures(tmp_path):
    out = _identity_roundtrip(make_pptx(tmp_path / "a.pptx"), tmp_path)
    shapes = pptx.Presentation(out).slides[0].shapes
    assert shapes.title.text == "Meningioma grading"
    out = _identity_roundtrip(make_docx(tmp_path / "a.docx"), tmp_path)
    assert docx.Document(out).paragraphs[0].text == "Meningioma grading"


@pytest.mark.private
def test_identity_roundtrip_private_corpus(tmp_path):
    listing = REPO_ROOT / "private" / "corpus.txt"
    paths = [Path(line.strip()) for line in listing.read_text(encoding="utf-8").splitlines() if line.strip()]
    done = 0
    for n, path in enumerate(p for p in paths if p.suffix.lower() in (".docx", ".pptx")):
        sub = tmp_path / str(n)
        sub.mkdir()
        out = _identity_roundtrip(path, sub)
        (docx.Document if out.suffix == ".docx" else pptx.Presentation)(out)
        done += 1
    assert done, "private/corpus.txt lists no DOCX / PPTX file"
