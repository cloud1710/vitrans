import zipfile

import docx
import pytest

from vitranslate.ooxml.package import OoxmlPackage, UnsupportedFile


def _make_docx(path):
    d = docx.Document()
    d.add_paragraph("Hello world")
    d.save(path)


def test_kind_and_byte_identical_save(tmp_path):
    src = tmp_path / "a.docx"
    _make_docx(src)
    pkg = OoxmlPackage(src)
    assert pkg.kind == "docx"
    root = pkg.root("word/document.xml")
    assert root is pkg.root("word/document.xml")  # parsed once
    out = tmp_path / "b.docx"
    pkg.save(out)
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(out) as b:
        assert [i.filename for i in a.infolist()] == [i.filename for i in b.infolist()]
        for name in a.namelist():
            assert a.read(name) == b.read(name), name  # nothing dirty -> byte identical


def test_dirty_part_is_reserialized(tmp_path):
    src = tmp_path / "a.docx"
    _make_docx(src)
    pkg = OoxmlPackage(src)
    root = pkg.root("word/document.xml")
    for t in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
        t.text = "Xin chao"
    pkg.mark_dirty("word/document.xml")
    out = tmp_path / "b.docx"
    pkg.save(out)
    assert docx.Document(out).paragraphs[0].text == "Xin chao"


def test_unsupported(tmp_path):
    bad = tmp_path / "x.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(UnsupportedFile):
        OoxmlPackage(bad)


def test_word_style_bytes_unchanged_without_dirty(tmp_path):
    src = tmp_path / "a.docx"
    _make_docx(src)
    word = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
            b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            b'<w:body><w:p><w:r><w:t>Hi</w:t></w:r></w:p></w:body></w:document>')
    word_src = tmp_path / "w.docx"
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(word_src, "w") as w:
        for i in a.infolist():
            data = word if i.filename == "word/document.xml" else a.read(i.filename)
            zi = zipfile.ZipInfo(i.filename, date_time=i.date_time)
            zi.compress_type = zipfile.ZIP_STORED if i.filename == "[Content_Types].xml" else zipfile.ZIP_DEFLATED
            w.writestr(zi, data)
    pkg = OoxmlPackage(word_src)
    pkg.root("word/document.xml")
    out = tmp_path / "o.docx"
    pkg.save(out)
    with zipfile.ZipFile(word_src) as a, zipfile.ZipFile(out) as b:
        for i in a.infolist():
            assert a.read(i.filename) == b.read(i.filename), i.filename
            assert i.compress_type == b.getinfo(i.filename).compress_type, i.filename


def test_mark_dirty_requires_parsed_root(tmp_path):
    src = tmp_path / "a.docx"
    _make_docx(src)
    pkg = OoxmlPackage(src)
    with pytest.raises(ValueError):
        pkg.mark_dirty("word/document.xml")


def test_save_is_atomic(tmp_path, monkeypatch):
    src = tmp_path / "a.docx"
    _make_docx(src)
    pkg = OoxmlPackage(src)
    pkg.root("word/document.xml")
    pkg.mark_dirty("word/document.xml")

    def boom(*a, **k):
        raise RuntimeError("boom")

    monkeypatch.setattr("vitranslate.ooxml.package.etree.tostring", boom)
    out_dir = tmp_path / "outdir"
    out_dir.mkdir()
    with pytest.raises(RuntimeError):
        pkg.save(out_dir / "b.docx")
    assert list(out_dir.iterdir()) == []


def test_save_keeps_comment_before_root(tmp_path):
    src = tmp_path / "a.docx"
    _make_docx(src)
    xml = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<!-- keep -->\n'
           b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    s2 = tmp_path / "c.docx"
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(s2, "w") as w:
        for i in a.infolist():
            w.writestr(i, xml if i.filename == "word/document.xml" else a.read(i.filename))
    pkg = OoxmlPackage(s2)
    pkg.root("word/document.xml")
    pkg.mark_dirty("word/document.xml")
    out = tmp_path / "o.docx"
    pkg.save(out)
    with zipfile.ZipFile(out) as z:
        assert b"<!-- keep -->" in z.read("word/document.xml")


def test_corrupt_deflate_data_is_unsupported(tmp_path):
    src = tmp_path / "a.docx"
    _make_docx(src)
    with zipfile.ZipFile(src) as z:
        info = z.getinfo("word/document.xml")
        assert info.compress_type == zipfile.ZIP_DEFLATED
        offset = info.header_offset
    data = bytearray(src.read_bytes())
    nlen = int.from_bytes(data[offset + 26:offset + 28], "little")
    elen = int.from_bytes(data[offset + 28:offset + 30], "little")
    start = offset + 30 + nlen + elen
    for k in range(start, start + 8):
        data[k] ^= 0xFF
    bad = tmp_path / "bad.docx"
    bad.write_bytes(bytes(data))
    with pytest.raises(UnsupportedFile):
        OoxmlPackage(bad)


def _declare_size(path, name, size):
    """Rewrite the uncompressed size of `name` in the central directory (a zip-bomb header; no huge data)."""
    data = bytearray(path.read_bytes())
    pos = data.find(b"PK\x01\x02")
    while pos != -1:
        nlen = int.from_bytes(data[pos + 28:pos + 30], "little")
        if bytes(data[pos + 46:pos + 46 + nlen]) == name.encode():
            data[pos + 24:pos + 28] = size.to_bytes(4, "little")
            break
        pos = data.find(b"PK\x01\x02", pos + 4)
    path.write_bytes(bytes(data))


def test_entry_declared_over_512mb_is_refused(tmp_path):
    src = tmp_path / "bomb.docx"
    _make_docx(src)
    _declare_size(src, "word/document.xml", 600 * 1024 * 1024)
    with pytest.raises(UnsupportedFile, match="too large"):
        OoxmlPackage(src)


def test_total_over_2gb_is_refused(tmp_path):
    src = tmp_path / "bomb2.docx"
    _make_docx(src)
    with zipfile.ZipFile(src) as z:
        names = [i.filename for i in z.infolist()][:5]
    for n in names:
        _declare_size(src, n, 500 * 1024 * 1024)  # each under the entry cap, sum over 2 GB
    with pytest.raises(UnsupportedFile, match="too large"):
        OoxmlPackage(src)
