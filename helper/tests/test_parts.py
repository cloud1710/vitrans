import docx
import pptx
from lxml import etree

from vitranslate.ooxml.ns import A, P
from vitranslate.ooxml.package import OoxmlPackage
from vitranslate.ooxml.parts import keep_paragraph, location, paragraph_kind, part_labels, select_parts


class FakePkg:
    def __init__(self, kind, names):
        self.kind = kind
        self._names = names

    def names(self):
        return self._names


def test_pptx_order_and_options():
    pkg = FakePkg("pptx", [
        "ppt/slides/slide10.xml", "ppt/slides/slide2.xml", "ppt/notesSlides/notesSlide2.xml",
        "ppt/charts/chart1.xml", "ppt/slideLayouts/slideLayout1.xml", "ppt/theme/theme1.xml",
        "ppt/presentation.xml",
    ])
    assert select_parts(pkg) == [
        "ppt/slides/slide2.xml", "ppt/slides/slide10.xml", "ppt/notesSlides/notesSlide2.xml",
        "ppt/charts/chart1.xml", "ppt/slideLayouts/slideLayout1.xml",
    ]
    assert "ppt/notesSlides/notesSlide2.xml" not in select_parts(pkg, notes=False)


def test_docx_parts_and_kinds():
    pkg = FakePkg("docx", [
        "word/document.xml", "word/header2.xml", "word/footer1.xml", "word/footnotes.xml",
        "word/comments.xml", "word/styles.xml", "word/charts/chart1.xml",
    ])
    assert select_parts(pkg) == [
        "word/document.xml", "word/header2.xml", "word/footer1.xml", "word/footnotes.xml",
        "word/charts/chart1.xml",
    ]
    assert "word/comments.xml" in select_parts(pkg, comments=True)
    assert paragraph_kind("docx", "word/document.xml") == "w"
    assert paragraph_kind("docx", "word/charts/chart1.xml") == "a"
    assert paragraph_kind("pptx", "ppt/slides/slide1.xml") == "a"


def test_locations():
    assert location("ppt/slides/slide3.xml") == "Slide 3"
    assert location("ppt/notesSlides/notesSlide3.xml") == "Notes 3"
    assert location("word/document.xml") == "Body"
    assert location("word/header1.xml") == "Header 1"
    assert location("word/footnotes.xml") == "Footnotes"


_NS = 'xmlns:a="%s" xmlns:p="%s"' % (A, P)


def _paras(xml):
    return etree.fromstring(xml).findall(".//{%s}p" % A)


_PH = ('<p:sp><p:nvSpPr><p:cNvPr id="1" name="x"/><p:cNvSpPr/><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>'
       '<p:txBody><a:p><a:r><a:t>Click</a:t></a:r></a:p></p:txBody></p:sp>')
_BOX = ('<p:sp><p:nvSpPr><p:cNvPr id="2" name="t"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        '<p:txBody><a:p><a:r><a:t>Box</a:t></a:r></a:p></p:txBody></p:sp>')


def _tree(shapes):
    return '<p:sld %s><p:cSld><p:spTree>%s</p:spTree></p:cSld></p:sld>' % (_NS, shapes)


def test_keep_paragraph():
    layout = "ppt/slideLayouts/slideLayout1.xml"
    assert keep_paragraph(layout, _paras(_tree(_PH))[0]) is False
    assert keep_paragraph(layout, _paras(_tree(_BOX))[0]) is True
    assert keep_paragraph(layout, _paras(_tree("<p:grpSp>%s</p:grpSp>" % _BOX))[0]) is True
    assert keep_paragraph("ppt/slides/slide1.xml", _paras(_tree(_PH))[0]) is True


def test_more_included_and_skipped_parts():
    pkg = FakePkg("pptx", [
        "ppt/slideMasters/slideMaster1.xml", "ppt/diagrams/data1.xml", "ppt/diagrams/drawing1.xml",
        "ppt/theme/theme1.xml", "ppt/notesMasters/notesMaster1.xml", "ppt/handoutMasters/handoutMaster1.xml",
        "ppt/presProps.xml", "ppt/diagrams/layout1.xml", "ppt/slides/slide1.xml.bak",
    ])
    assert select_parts(pkg) == [
        "ppt/diagrams/data1.xml", "ppt/diagrams/drawing1.xml", "ppt/slideMasters/slideMaster1.xml",
    ]
    d = FakePkg("docx", [
        "word/header.xml", "word/endnotes.xml", "word/diagrams/data1.xml", "word/diagrams/drawing1.xml",
        "word/glossary/document.xml", "word/theme/theme1.xml",
    ])
    assert select_parts(d) == [
        "word/header.xml", "word/endnotes.xml", "word/diagrams/data1.xml", "word/diagrams/drawing1.xml",
    ]
    assert paragraph_kind("docx", "word/diagrams/data1.xml") == "a"


def _deck(tmp_path):
    prs = pptx.Presentation()
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "T%d" % (i + 1)
    prs.slides[2].notes_slide.notes_text_frame.text = "note"
    path = tmp_path / "d.pptx"
    prs.save(path)
    return path


def test_labels_used_layouts_and_slide_order(tmp_path):
    path = _deck(tmp_path)
    pkg = OoxmlPackage(path)
    labels = part_labels(pkg)
    assert labels["ppt/slides/slide3.xml"] == "Slide 3"
    assert labels["ppt/notesSlides/notesSlide1.xml"] == "Notes 3"
    sel = select_parts(pkg)
    layouts = [n for n in sel if "slideLayouts" in n]
    assert layouts == ["ppt/slideLayouts/slideLayout2.xml"]
    assert [n for n in sel if "slideMasters" in n] == ["ppt/slideMasters/slideMaster1.xml"]
    assert [n for n in sel if "/slides/" in n] == [
        "ppt/slides/slide1.xml", "ppt/slides/slide2.xml", "ppt/slides/slide3.xml"]
    # move slide 3 to the first position
    root = pkg.root("ppt/presentation.xml")
    lst = root.find(".//{%s}sldIdLst" % P)
    last = lst[-1]
    lst.remove(last)
    lst.insert(0, last)
    pkg.mark_dirty("ppt/presentation.xml")
    out = tmp_path / "e.pptx"
    pkg.save(out)
    pkg2 = OoxmlPackage(out)
    labels2 = part_labels(pkg2)
    assert labels2["ppt/slides/slide3.xml"] == "Slide 1"
    assert labels2["ppt/slides/slide1.xml"] == "Slide 2"
    assert labels2["ppt/notesSlides/notesSlide1.xml"] == "Notes 1"
    sel2 = select_parts(pkg2)
    assert [n for n in sel2 if "/slides/" in n][0] == "ppt/slides/slide3.xml"
    assert sel2.index("ppt/notesSlides/notesSlide1.xml") > sel2.index("ppt/slides/slide2.xml")


def test_label_fallback_for_docx(tmp_path):
    p = tmp_path / "a.docx"
    docx.Document().save(p)
    assert part_labels(OoxmlPackage(p))["word/document.xml"] == "Body"


def _save_with_sldids(path, out, keep):
    pkg = OoxmlPackage(path)
    root = pkg.root("ppt/presentation.xml")
    lst = root.find(".//{%s}sldIdLst" % P)
    for i, el in enumerate(list(lst)):
        if i not in keep:
            lst.remove(el)
    pkg.mark_dirty("ppt/presentation.xml")
    pkg.save(out)
    return OoxmlPackage(out)


def test_orphan_slide_and_its_notes_skipped(tmp_path):
    path = _deck(tmp_path)
    pkg = _save_with_sldids(path, tmp_path / "o.pptx", keep={0, 1})  # slide 3 (with notes) dropped
    sel = select_parts(pkg)
    assert "ppt/slides/slide3.xml" not in sel
    assert not [n for n in sel if "notesSlides" in n]
    labels = part_labels(pkg)
    assert labels["ppt/slides/slide1.xml"] == "Slide 1"
    assert labels["ppt/slides/slide2.xml"] == "Slide 2"


def test_notes_ordered_by_slide_number(tmp_path):
    prs = pptx.Presentation()
    for _ in range(3):
        prs.slides.add_slide(prs.slide_layouts[1])
    prs.slides[2].notes_slide.notes_text_frame.text = "n3"  # becomes notesSlide1
    prs.slides[0].notes_slide.notes_text_frame.text = "n1"  # becomes notesSlide2
    path = tmp_path / "n.pptx"
    prs.save(path)
    pkg = OoxmlPackage(path)
    labels = part_labels(pkg)
    assert labels["ppt/notesSlides/notesSlide1.xml"] == "Notes 3"
    assert labels["ppt/notesSlides/notesSlide2.xml"] == "Notes 1"
    notes = [n for n in select_parts(pkg) if "notesSlides" in n]
    assert notes == ["ppt/notesSlides/notesSlide2.xml", "ppt/notesSlides/notesSlide1.xml"]
