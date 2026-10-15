import json
import shutil

import docx
import pptx
import pytest

from fixtures import make_docx, make_pptx
from nonmedical import BUDGET_MEMO
from packutil import BREAST, build, copy_real, make_pack
from vitranslate import config, pipeline
from vitranslate.device import Device
from vitranslate.llm import FakeLLM, LLMError
from vitranslate.matching import THRESHOLD
from vitranslate.ooxml.document import collect_items
from vitranslate.ooxml.package import OoxmlPackage, UnsupportedFile
from vitranslate.packs import load_pack
from vitranslate.pipeline import Options, output_name, run_job
from vitranslate.registry import Registry
from vitranslate.tm import TM

VI = {"Meningioma grading": "Ph\u00e2n \u0111\u1ed9 u m\u00e0ng n\u00e3o",
      "Negative staining": "Nhu\u1ed9m \u00e2m t\u00ednh",
      # a real translation: the default fake ("Speaker note about ph\u00e2n \u0111\u1ed9") copies three English words,
      # which the copied-English gate rightly flags as english_left (and the read-back pass then skips the line)
      "Speaker note about grading": "Ghi ch\u00fa c\u1ee7a ng\u01b0\u1eddi tr\u00ecnh b\u00e0y v\u1ec1 ph\u00e2n \u0111\u1ed9"}


def fake_translate(system, user):
    text = user.split("Translate:\n", 1)[1]
    return VI.get(text, text.replace("grading", "ph\u00e2n \u0111\u1ed9"))


class AllFonts:
    def covers(self, family, chars):
        return True

    def classify(self, family):
        return "sans"


class ClosingLLM(FakeLLM):
    def __init__(self, **kw):
        super().__init__(**kw)
        self.closed = False

    def close(self):
        self.closed = True


@pytest.fixture(autouse=True)
def registry_env(tmp_path, monkeypatch, pathology_registry):
    """Default registry of run_job = a folder built from packs/pathology; packs unpack into a per-test app home."""
    monkeypatch.setenv("VITRANSLATE_REGISTRY", str(pathology_registry))
    monkeypatch.setenv("VITRANSLATE_HOME", str(tmp_path / "apphome"))


def run(path, tmp_path, events=None, llm=None, sleep=None, options=None, **kw):
    if "tm" not in kw:
        kw["tm"] = TM(tmp_path / "tm.sqlite")
    if sleep is not None:
        kw["sleep"] = sleep
    llm = llm or FakeLLM(fn=fake_translate)
    return run_job(path, tmp_path / "out", options or Options(),
                   progress=(events if events is not None else []).append,
                   llm_factory=lambda dev, model: llm, probe=kw.pop("probe", lambda: Device("cpu", "test cpu", 0, 0)),
                   font_index=AllFonts(), **kw)


def memo_docx(path):
    d = docx.Document()
    for line in BUDGET_MEMO.split("\n"):
        d.add_paragraph(line)
    d.save(path)
    return path


def files(d):
    return sorted(p.name for p in d.iterdir()) if d.exists() else []


def test_output_name(tmp_path):
    assert output_name(tmp_path / "a b.pptx") == "a b_vi.pptx"


def test_pptx_job(tmp_path):
    events = []
    llm = ClosingLLM(fn=fake_translate)
    res = run(make_pptx(tmp_path / "a.pptx"), tmp_path, events, llm=llm)
    assert res.output.name == "a_vi.pptx"
    assert files(tmp_path / "out") == ["a_vi.pptx"]  # no temporary file left behind
    prs = pptx.Presentation(res.output)
    assert prs.slides[0].shapes.title.text == VI["Meningioma grading"]
    assert prs.slides[1].shapes.title.text == "References"
    rep = res.report
    assert rep["domain"] == "pathology" and rep["device"] == "cpu" and rep["model"] == "qwen3.5:9b"
    assert rep["domain_name"] == load_pack(config.REPO_ROOT / "packs" / "pathology").name
    assert rep["counts"]["segments"] >= 8 and rep["skipped"]["references"] >= 2
    stages = [e["stage"] for e in events]
    assert stages[0] == "extract" and stages[-1] == "done" and "translate" in stages
    assert events[stages.index("translate")]["total"] == rep["counts"]["segments"]
    assert llm.closed


def test_docx_job_reports_tracked_changes(tmp_path):
    res = run(make_docx(tmp_path / "a.docx"), tmp_path, [])
    assert docx.Document(res.output).paragraphs[0].text == VI["Meningioma grading"]
    assert any(s["reason"] == "tracked changes" for s in res.report["skipped_items"])


@pytest.mark.parametrize("maker,name", [(make_pptx, "a.pptx"), (make_docx, "a.docx")])
def test_every_item_is_translated_or_skipped(tmp_path, maker, name):
    path = maker(tmp_path / name)
    items = collect_items(OoxmlPackage(path), notes=True, comments=False)
    rep = run(path, tmp_path).report
    assert rep["counts"]["segments"] + sum(rep["skipped"].values()) == len(items)
    assert len(rep["skipped_items"]) == sum(n for reason, n in rep["skipped"].items() if reason != "no_text")


def test_extension_must_match_content(tmp_path):
    path = make_docx(tmp_path / "a.pptx")  # a Word document renamed to .pptx
    with pytest.raises(UnsupportedFile) as exc:
        run(path, tmp_path)
    msg = str(exc.value)
    assert ".pptx" in msg and str(tmp_path) not in msg and "a.pptx" not in msg
    assert files(tmp_path / "out") == []


def test_ollama_failing_twice_aborts_the_job(tmp_path, home, monkeypatch):
    opened = []

    class RecTM(TM):
        def __init__(self, path):
            super().__init__(path)
            self.closed = False
            opened.append(self)

        def close(self):
            self.closed = True
            super().close()

    def down(system, user):
        raise LLMError("connection refused")

    monkeypatch.setattr(pipeline, "TM", RecTM)
    llm, slept = ClosingLLM(fn=down), []
    with pytest.raises(LLMError) as exc:
        run(make_docx(tmp_path / "a.docx"), tmp_path, llm=llm, sleep=slept.append, tm=None)
    assert "twice" in str(exc.value) and "connection refused" in str(exc.value)
    assert slept == [3] and len(llm.calls) == 2
    assert files(tmp_path / "out") == []
    assert len(opened) == 1 and opened[0].closed and llm.closed


def test_ollama_failing_once_is_retried(tmp_path):
    calls = []

    def flaky(system, user):
        calls.append(user)
        if len(calls) == 1:
            raise LLMError("timed out")
        return fake_translate(system, user)

    slept = []
    res = run(make_docx(tmp_path / "a.docx"), tmp_path, llm=FakeLLM(fn=flaky), sleep=slept.append)
    assert slept == [3]
    assert docx.Document(res.output).paragraphs[0].text == VI["Meningioma grading"]


def test_output_that_does_not_reopen_is_discarded(tmp_path, monkeypatch):
    def broken(path):
        raise ValueError("cannot read " + str(path))

    monkeypatch.setattr(pipeline, "reopen", broken)
    with pytest.raises(pipeline.JobError) as exc:
        run(make_pptx(tmp_path / "a.pptx"), tmp_path)
    assert exc.value.report["counts"]["segments"] >= 8 and exc.value.report["output"] == "a_vi.pptx"
    assert str(tmp_path) not in str(exc.value)
    assert files(tmp_path / "out") == []


REAL_META = json.loads((config.REPO_ROOT / "packs" / "pathology" / "pack.json").read_text(encoding="utf-8"))
VERSION = REAL_META["version"]


def classify_event(events):
    return next(e for e in events if e["stage"] == "classify")


def test_pack_selected_downloaded_then_cached(tmp_path):
    events = []
    rep = run(make_pptx(tmp_path / "a.pptx"), tmp_path, events).report
    (p,) = rep["packs"]
    assert p["id"] == "pathology" and p["version"] == VERSION and p["source"] == "downloaded"
    assert p["reason"] == "match" and p["score"] >= THRESHOLD and "meningioma" in p["matched"]
    assert len(p["matched"]) <= 8 and p["name"] == rep["domain_name"] == REAL_META["name"]
    assert rep["packs_unavailable"] == [] and rep["packs_unavailable_names"] == []
    assert rep["pack"] == "pathology@" + VERSION and rep["registry_offline"] is False
    cls = classify_event(events)
    assert [x["id"] for x in cls["packs"]] == ["pathology"] and cls["offline"] is False and cls["unavailable"] == []
    assert set(cls["packs"][0]) == {"id", "name", "score", "matched", "reason"}
    assert (tmp_path / "apphome" / "packs" / ("pathology@" + VERSION) / "glossary.csv").is_file()
    rep2 = run(make_pptx(tmp_path / "b.pptx"), tmp_path).report
    assert [(x["id"], x["source"]) for x in rep2["packs"]] == [("pathology", "cache")]


def test_matched_terms_are_shown_as_written_in_the_pack(tmp_path):
    d = docx.Document()
    d.add_paragraph("Immunohistochemistry (IHC) staining of the tumor biopsy; WHO classification grade 2; necrosis.")
    d.save(tmp_path / "ihc.docx")
    events = []
    (p,) = run(tmp_path / "ihc.docx", tmp_path, events).report["packs"]
    assert "IHC" in p["matched"] and "WHO classification" in p["matched"] and "ihc" not in p["matched"]
    assert classify_event(events)["packs"][0]["matched"] == p["matched"]


def test_tm_key_carries_each_pack_sha(tmp_path, monkeypatch, pathology_registry):
    keys = []

    class Recording(pipeline.Translator):
        def __init__(self, llm, system, model, pack_key, *args, **kw):
            keys.append(pack_key)
            super().__init__(llm, system, model, pack_key, *args, **kw)

    monkeypatch.setattr(pipeline, "Translator", Recording)
    run(make_pptx(tmp_path / "a.pptx"), tmp_path)
    manifest = json.loads((pathology_registry / "manifest.json").read_text(encoding="utf-8"))
    sha = manifest["packs"][0]["sha256"]
    assert keys[0].startswith("pathology@%s#%s+_style@" % (VERSION, sha[:8]))
    assert keys[0].endswith("+" + config.VERSION)


def test_non_medical_docx_runs_without_a_pack(tmp_path):
    llm = FakeLLM(fn=fake_translate)
    rep = run(memo_docx(tmp_path / "memo.docx"), tmp_path, llm=llm, options=Options(domain="")).report
    assert rep["domain"] == "_none" and rep["domain_name"] == "Chung" and rep["packs"] == []
    assert rep["packs_unavailable"] == [] and rep["pack"] == "_none@0"
    assert llm.calls and not any("Glossary" in system for system, _ in llm.calls)
    assert not (tmp_path / "apphome" / "packs").exists()  # nothing downloaded


def test_unknown_or_broken_packs_are_reported_unavailable(tmp_path):
    rep = run(make_pptx(tmp_path / "a.pptx"), tmp_path, options=Options(domain="pathology,no_such")).report
    assert [p["id"] for p in rep["packs"]] == ["pathology"] and rep["packs_unavailable"] == ["no_such"]
    assert rep["packs"][0]["reason"] == "user" and rep["packs_unavailable_names"] == ["no_such"]
    rep = run(make_pptx(tmp_path / "b.pptx"), tmp_path, options=Options(domain="no_such")).report
    assert rep["domain"] == "_none" and rep["packs"] == [] and rep["packs_unavailable"] == ["no_such"]
    # a pack whose download fails its sha256 check is skipped, the job still runs
    reg_dir = tmp_path / "broken"
    build(copy_real(tmp_path / "src", "pathology"), reg_dir)
    zf = next(reg_dir.glob("pathology-*.zip"))
    zf.write_bytes(zf.read_bytes()[:-10])
    events = []
    reg = Registry(reg_dir, tmp_path / "home2")
    rep = run(make_pptx(tmp_path / "c.pptx"), tmp_path, events, registry=reg).report
    assert rep["domain"] == "_none" and rep["packs"] == [] and rep["packs_unavailable"] == ["pathology"]
    assert rep["packs_unavailable_names"] == [REAL_META["name"]]
    assert classify_event(events)["packs"] == [] and classify_event(events)["unavailable"] == [REAL_META["name"]]


def test_child_pack_brings_its_parent(tmp_path):
    src = copy_real(tmp_path / "src", "pathology")
    make_pack(src, "breast", parent="pathology", **BREAST)
    build(src, tmp_path / "reg")
    reg = Registry(tmp_path / "reg", tmp_path / "home2")
    rep = run(make_pptx(tmp_path / "a.pptx"), tmp_path, registry=reg, options=Options(domain="breast")).report
    assert [(p["id"], p["reason"]) for p in rep["packs"]] == [("breast", "user"), ("pathology", "parent")]
    assert rep["domain"] == "breast+pathology"
    # the user names the parent first: the child still merges first
    rep = run(make_pptx(tmp_path / "b.pptx"), tmp_path, registry=reg,
              options=Options(domain="pathology,breast")).report
    assert [(p["id"], p["reason"]) for p in rep["packs"]] == [("breast", "user"), ("pathology", "user")]


MAT_SANG = "m" + chr(0x1EB7) + "t s" + chr(0xE0) + "ng"   # "mat sang" with Vietnamese diacritics
DANG_SANG = "d" + chr(0x1EA1) + "ng s" + chr(0xE0) + "ng"  # "dang sang", the pathology value
# pathology outscores breast on this text; "cribriform" is a pathology row that the breast pack overrides
PARENT_HIGHER = (
    "Histopathology of the biopsy showed a meningioma with brisk mitoses and focal necrosis. Immunohistochemistry "
    "(IHC) staining was negative for the tumor markers; WHO classification grade 2. No lymph node metastasis. "
    "Glioma, astrocytoma and sarcoma with necrosis, gliosis and dementia; carcinoma and lymphoma on cytology; "
    "neoplasm with metastases. HER2 and mammography; BI-RADS. A cribriform pattern.")


def test_child_override_wins_when_the_parent_scores_higher(tmp_path):
    src = copy_real(tmp_path / "src", "pathology")
    breast = dict(BREAST, glossary=BREAST["glossary"] + [("cribriform", MAT_SANG)])
    make_pack(src, "breast", parent="pathology", **breast)
    build(src, tmp_path / "reg")
    d = docx.Document()
    d.add_paragraph(PARENT_HIGHER)
    d.save(tmp_path / "deck.docx")
    llm = FakeLLM(fn=fake_translate)
    rep = run(tmp_path / "deck.docx", tmp_path, llm=llm, registry=Registry(tmp_path / "reg", tmp_path / "home2")).report
    assert [p["id"] for p in rep["packs"]] == ["breast", "pathology"]
    assert rep["packs"][1]["score"] > rep["packs"][0]["score"]  # the parent scores higher, the child still wins
    system = llm.calls[0][0]
    assert "- cribriform = " + MAT_SANG in system and DANG_SANG not in system


def test_pack_that_fails_to_load_is_unavailable_and_forgotten(tmp_path):
    src = tmp_path / "src"
    make_pack(src, "bad", detect=["tumor"])
    (src / "bad" / "rules.json").write_text(json.dumps({"term_fixes": [{"en": "x", "vi": "y", "wrong": ["(oops"]}]}),
                                            encoding="utf-8")
    build(src, tmp_path / "reg")
    reg = Registry(tmp_path / "reg", tmp_path / "home2")
    events = []
    rep = run(make_pptx(tmp_path / "a.pptx"), tmp_path, events, registry=reg, options=Options(domain="bad")).report
    assert rep["packs"] == [] and rep["packs_unavailable"] == ["bad"] and rep["packs_unavailable_names"] == ["Bad"]
    assert rep["domain"] == "_none" and classify_event(events)["unavailable"] == ["Bad"]
    assert not (tmp_path / "home2" / "packs" / "bad@1.0.0" / ".sha256").exists()  # downloaded again next time


def test_bad_manifest_entries_never_crash_a_job(tmp_path, pathology_registry):
    reg_dir = tmp_path / "reg"
    shutil.copytree(pathology_registry, reg_dir)
    m = json.loads((reg_dir / "manifest.json").read_text(encoding="utf-8"))
    m["packs"].insert(0, {"id": "broken", "profile": {"tumor": "x"}, "detect": ["tumor"]})
    m["packs"].append("junk")
    m["idf"]["tumor"] = None
    (reg_dir / "manifest.json").write_text(json.dumps(m), encoding="utf-8")
    rep = run(make_pptx(tmp_path / "a.pptx"), tmp_path, registry=Registry(reg_dir, tmp_path / "home2")).report
    assert [p["id"] for p in rep["packs"]] == ["pathology"]


def test_offline_registry_without_cache_translates_without_packs(tmp_path, monkeypatch):
    monkeypatch.setenv("VITRANSLATE_REGISTRY", str(tmp_path / "missing"))
    events = []
    rep = run(make_pptx(tmp_path / "a.pptx"), tmp_path, events).report
    assert rep["domain"] == "_none" and rep["packs"] == [] and rep["registry_offline"] is True
    assert classify_event(events) == {"stage": "classify", "packs": [], "unavailable": [], "offline": True,
                                      "candidates": []}


GPU = Device("gpu", "test gpu", 16000, 999)
TITLE_EDITED = "Ph\u00e2n \u0111\u1ed9 c\u1ee7a u m\u00e0ng n\u00e3o"   # the editor's version of VI["Meningioma grading"] (plumbing only)
NEG_EDITED = "Nhu\u1ed9m \u00e2m t\u00ednh 2"  # a number the draft does not have: rejected


def translate_then_edit(system, user):
    """The translator fake, and an editor that rewrites two known drafts and returns every other line unchanged."""
    from vitranslate.polish import EDITOR_HEAD
    if not system.startswith(EDITOR_HEAD.split("(")[0]):
        return fake_translate(system, user)
    out, n = [], 0
    for line in user.splitlines():
        if line.startswith("[") and "] EN: " in line:
            n = int(line[1:line.index("]")])
        elif line.startswith("VI: "):
            draft = line[4:]
            out.append("[%d] %s" % (n, {VI["Meningioma grading"]: TITLE_EDITED,
                                       VI["Negative staining"]: NEG_EDITED}.get(draft, draft)))
    return "\n".join(out)


def editor_calls(llm):
    from vitranslate.polish import EDITOR_HEAD
    return [u for s, u in llm.calls if s.startswith(EDITOR_HEAD.split("(")[0])]


def test_polish_edits_the_output_before_writing(tmp_path):
    events, llm = [], FakeLLM(fn=translate_then_edit)
    res = run(make_pptx(tmp_path / "a.pptx"), tmp_path, events, llm=llm, options=Options(polish=True))
    prs = pptx.Presentation(res.output)
    assert prs.slides[0].shapes.title.text == TITLE_EDITED
    rep = res.report
    pol = rep["polish"]
    assert pol["enabled"] and pol["changed"] == 1 and pol["rejected"] == {"numbers": 1}
    assert pol["groups"] >= 2 and pol["lines"] >= 2 and pol["calls"] == len(editor_calls(llm))
    assert rep["info"]["polish_rejected"] == 1 and "polish" in rep["seconds"]
    stages = [e["stage"] for e in events]
    assert stages.index("translate") < stages.index("polish") < stages.index("write")
    last = [e for e in events if e["stage"] == "polish"][-1]
    assert last["done"] == last["total"] == pol["groups"]


def test_polish_off_or_auto_on_cpu_never_calls_the_editor(tmp_path):
    for options in (Options(polish=False), Options()):
        llm = FakeLLM(fn=translate_then_edit)
        rep = run(make_pptx(tmp_path / "a.pptx"), tmp_path, llm=llm, options=options).report
        assert editor_calls(llm) == [] and rep["polish"] == {"enabled": False}
    llm = FakeLLM(fn=translate_then_edit)
    rep = run(make_pptx(tmp_path / "b.pptx"), tmp_path, llm=llm, probe=lambda: GPU).report
    assert editor_calls(llm) and rep["polish"]["enabled"]  # auto: on with a GPU
    llm = FakeLLM(fn=translate_then_edit)
    run(make_pptx(tmp_path / "c.pptx"), tmp_path, llm=llm, probe=lambda: GPU, options=Options(polish=False))
    assert editor_calls(llm) == []


def test_progress_events_carry_counts_locations_and_unselected_candidates(tmp_path):
    src = copy_real(tmp_path / "src", "pathology")
    for pid, words in (("derm", ["melanoma", "nevus", "epidermis"]), ("cardio", ["myocardium", "infarct", "valve"]),
                       ("gastro", ["gastritis", "colon", "polyp"]), ("renal", ["glomerulus", "nephron", "tubule"])):
        make_pack(src, pid, detect=words, glossary=[(w, "x " + w) for w in words])
    build(src, tmp_path / "reg")
    events = []
    rep = run(make_pptx(tmp_path / "a.pptx"), tmp_path, events,
              registry=Registry(tmp_path / "reg", tmp_path / "home2")).report
    counts = [e for e in events if e["stage"] == "extract" and "segments" in e]
    assert len(counts) == 1 and counts[0]["segments"] == rep["counts"]["segments"] and counts[0]["words"] > 0
    translate = [e for e in events if e["stage"] == "translate"]
    assert all(e["where"].startswith(("Slide", "Notes")) for e in translate)
    ev = classify_event(events)
    assert [p["id"] for p in ev["packs"]] == ["pathology"] and ev["packs"][0]["reason"] == "match"
    cands = ev["candidates"]
    assert len(cands) == 3 and all(set(c) == {"id", "name", "score"} for c in cands)
    assert "pathology" not in [c["id"] for c in cands]
    assert [c["score"] for c in cands] == sorted([c["score"] for c in cands], reverse=True)


def test_estimate_in_the_device_event_and_a_refined_eta_while_translating(tmp_path):
    from vitranslate.recommend import job_estimate, rates
    seen = []

    def estimator(dev, model, words, polish):
        seen.append((dev.kind, model, list(words), polish))
        return job_estimate(rates(dev.kind, model, {}), dev.kind, words, polish)
    events = []
    rep = run(make_pptx(tmp_path / "a.pptx"), tmp_path, events, estimator=estimator).report
    kind, model, words, polish = seen[0]
    assert (kind, model, polish) == ("cpu", config.MODEL_DEFAULT_CPU, False) and len(words) == rep["counts"]["segments"]
    dev = next(e for e in events if e["stage"] == "device")
    assert dev["polish"] is False and set(dev["estimate"]) == {"translate_s", "polish_s", "write_s", "total_s", "basis"}
    assert dev["estimate"]["polish_s"] == 0 and dev["estimate"]["total_s"] > 0
    etas = [e["eta_s"] for e in events if e["stage"] == "translate"]
    assert etas and all(isinstance(x, (int, float)) and x >= 0 for x in etas)
    # without an estimator (tools, tests) the events stay as they were
    events = []
    run(make_pptx(tmp_path / "b.pptx"), tmp_path, events)
    dev = next(e for e in events if e["stage"] == "device")
    assert "estimate" not in dev and not any("eta_s" in e for e in events)


@pytest.mark.parametrize("dev, fast, want", [
    (GPU, False, "gemma4:e4b"),  # spec Amendment E: Gemma 4 E4B with a GPU
    (Device("cpu", "test cpu", 0, 0), False, "qwen3.5:9b"),  # Qwen 3.5 9B on the CPU until Gemma is measured there
    (GPU, True, "gemma4:e2b"),  # the fast draft on either device
    (Device("cpu", "test cpu", 0, 0), True, "gemma4:e2b"),
])
def test_the_model_follows_the_device_and_the_fast_option(tmp_path, dev, fast, want):
    from vitranslate.recommend import job_estimate, rates
    seen = []

    def estimator(d, model, words, polish):
        seen.append(model)
        return job_estimate(rates(d.kind, model, {}), d.kind, words, polish)
    rep = run(make_pptx(tmp_path / "a.pptx"), tmp_path, probe=lambda: dev, options=Options(fast=fast, polish=False),
              estimator=estimator).report
    assert seen == [want] and rep["model"] == want


def test_report_records_the_fast_draft_and_the_review_level(tmp_path):
    rep = run(make_pptx(tmp_path / "a.pptx"), tmp_path, options=Options(fast=True)).report
    assert rep["fast"] is True and rep["model"] == config.MODEL_FAST
    assert rep["review"]["level"] == "thorough" and "fast_draft" in rep["review"]["reasons"]


LYM_SRC = "Many lymphocytes were seen in the biopsy of the tumor with necrosis."
LYM_HOUSE = "Nhi\u1ec1u lymph\u00f4 b\u00e0o trong sinh thi\u1ebft u c\u00f3 ho\u1ea1i t\u1eed."


def lym_docx(path):
    d = docx.Document()
    d.add_paragraph(LYM_SRC)
    d.save(path)
    return path


def test_house_forms_of_the_selected_pack_are_no_strange_words(tmp_path):
    # the pathology term-lock writes lymphocyte as "lymph\u00f4 b\u00e0o" (his form): the suspicious-word check must
    # accept it, and still flag a strange word no active pack writes
    answer = "Nhi\u1ec1u t\u1ebf b\u00e0o lympho trong sinh thi\u1ebft u c\u00f3 ho\u1ea1i t\u1eed, nghi carcin\u00f4m."
    res = run(lym_docx(tmp_path / "l.docx"), tmp_path, llm=FakeLLM(fn=lambda s, u: answer), options=Options(polish=False))
    rep = res.report
    assert [p["id"] for p in rep["packs"]] == ["pathology"]
    assert docx.Document(res.output).paragraphs[0].text == LYM_HOUSE[:-1] + ", nghi carcin\u00f4m."
    flags = [f for r in rep["flagged"] for f in r["flags"]]
    assert "term_fix:lymphocyte" in flags and "suspicious:carcin\u00f4m" in flags
    assert not any(f.startswith("suspicious:lymph") for f in flags)


def test_the_read_back_pass_may_write_a_house_form_of_the_selected_pack(tmp_path):
    from vitranslate.polish import EDITOR_HEAD
    draft = "Nhi\u1ec1u b\u1ea1ch c\u1ea7u lympho trong sinh thi\u1ebft u c\u00f3 ho\u1ea1i t\u1eed."

    def fn(system, user):
        return "[1] " + LYM_HOUSE if system.startswith(EDITOR_HEAD.split("(")[0]) else draft
    res = run(lym_docx(tmp_path / "l.docx"), tmp_path, llm=FakeLLM(fn=fn), options=Options(polish=True))
    assert res.report["polish"]["changed"] == 1 and res.report["polish"]["rejected"] == {}
    assert docx.Document(res.output).paragraphs[0].text == LYM_HOUSE
