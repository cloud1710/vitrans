"""Read-back editor pass (spec Amendment B): grouping, prompt, answer parsing, per-line guards, TM, pipeline."""
import docx
from pptx import Presentation
from pptx.util import Inches

from vitranslate.ooxml.document import collect_items
from vitranslate.ooxml.package import OoxmlPackage
from vitranslate.packs import Pack
from vitranslate.polish import (build_editor_system, build_editor_user, group_hits, group_lines, guard_line,
                                parse_answer)
from vitranslate.translate import Outcome

U = "u"
NEURON_VI = "Neuron ho\u1ea1t \u0111\u1ed9ng b\u00ecnh th\u01b0\u1eddng"  # draft of "Normal functioning neurons"
TUMOR_HIT = ("tumor", U)
NECROSIS_HIT = ("necrosis", "ho\u1ea1i t\u1eed")


def drafts(items, flags=None):
    """(item, outcome) pairs whose drafts are the sources with a "vi " prefix; flags: {index: [flags]}."""
    flags = flags or {}
    work = [i for i in items if i.segment]
    return [(it, Outcome("vi " + it.segment.tagged, it.segment, 0, list(flags.get(n, []))))
            for n, it in enumerate(work)]


def texts(pairs, groups):
    return [[pairs[i][0].segment.plain for i in g] for g in groups]


def deck(path):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Normal functioning neurons"
    s1.placeholders[1].text_frame.text = "Synthetic bullet line kept longer than the cap"
    s1.notes_slide.notes_text_frame.text = "Speaker note about the neurons"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Practical aspects of testing"
    table = s2.shapes.add_table(2, 1, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Marker panel"
    table.cell(1, 0).text = "Negative staining"
    prs.save(path)
    return path


def test_pptx_groups_are_slides_and_notes_pages(tmp_path):
    pairs = drafts(collect_items(OoxmlPackage(deck(tmp_path / "d.pptx"))))
    groups = group_lines(pairs, "pptx")
    assert texts(pairs, groups) == [
        ["Normal functioning neurons", "Synthetic bullet line kept longer than the cap"],
        ["Practical aspects of testing", "Marker panel", "Negative staining"],
        ["Speaker note about the neurons"]]  # notes parts come after the slides


def test_kept_flagged_locked_and_tag_broken_lines_are_never_sent(tmp_path):
    pairs = drafts(collect_items(OoxmlPackage(deck(tmp_path / "d.pptx"))),
                   {0: ["kept_code"], 1: ["kept_english"], 2: ["kept_reference"], 3: ["unchanged"],
                    4: ["english_left", "retried:english"]})
    pairs[5][1].text = "vi Speaker note</g1>"  # tags no longer exact against the source
    assert group_lines(pairs, "pptx") == []
    pairs = drafts(collect_items(OoxmlPackage(deck(tmp_path / "d.pptx"))), {0: ["empty_output"]})
    pairs[1][1].text = "vi line one\nline two"
    assert texts(pairs, group_lines(pairs, "pptx"))[0] == ["Practical aspects of testing", "Marker panel",
                                                          "Negative staining"]
    pairs = drafts(collect_items(OoxmlPackage(deck(tmp_path / "d.pptx"))), {0: ["tags_broken"]})
    pairs[1][0].segment.locked = ["field text"]
    assert texts(pairs, group_lines(pairs, "pptx")) == [
        ["Practical aspects of testing", "Marker panel", "Negative staining"], ["Speaker note about the neurons"]]


def test_a_slide_over_the_cap_is_split_in_order(tmp_path, monkeypatch):
    from vitranslate import polish
    pairs = drafts(collect_items(OoxmlPackage(deck(tmp_path / "d.pptx"))))
    monkeypatch.setattr(polish, "SLIDE_MAX_LINES", 2)
    assert [len(g) for g in group_lines(pairs, "pptx")] == [2, 2, 1, 1]
    monkeypatch.setattr(polish, "SLIDE_MAX_LINES", 20)
    monkeypatch.setattr(polish, "SLIDE_MAX_CHARS", 30)  # source characters: one long line still forms a group
    assert texts(pairs, group_lines(pairs, "pptx"))[:2] == [["Normal functioning neurons"],
                                                            ["Synthetic bullet line kept longer than the cap"]]


def passage_docx(path):
    d = docx.Document()
    d.add_heading("Background", 1)
    for n in range(10):
        d.add_paragraph("Body paragraph number %d of the background." % n)
    d.add_heading("Methods", 2)
    d.add_paragraph("We stained the sections.")
    t = d.add_table(rows=2, cols=1)
    t.cell(0, 0).text = "Marker"
    t.cell(1, 0).text = "Ki-67 index"
    d.add_paragraph("After the table.")
    d.add_section()
    d.add_paragraph("New section text.")
    d.save(path)
    return path


def test_docx_passages_break_at_headings_tables_sections_and_eight_lines(tmp_path):
    pairs = drafts(collect_items(OoxmlPackage(passage_docx(tmp_path / "p.docx"))))
    got = texts(pairs, group_lines(pairs, "docx"))
    body = ["Body paragraph number %d of the background." % n for n in range(10)]
    assert got == [["Background"] + body[:7], body[7:], ["Methods", "We stained the sections."],
                   ["Marker", "Ki-67 index"], ["After the table."], ["New section text."]]


def test_editor_system_carries_style_pack_voice_and_rules():
    pack = Pack("p", "1", "Pathology", [], "PACK VOICE", [], [], [])
    system = build_editor_system("STYLE RULES", pack)
    assert system.index("STYLE RULES") < system.index("PACK VOICE")
    for needle in ("noun phrase", "C\u00e1c", "Nh\u1eefng", "<g1>", "<x1/>", "[n]", "unchanged", "Pathology"):
        assert needle in system, needle


def test_editor_user_numbers_every_line_with_its_source_and_lists_the_group_glossary():
    lines = [("Normal <g1>functioning</g1> neurons", NEURON_VI), ("Tumor necrosis", "draft two")]
    user = build_editor_user(lines, [TUMOR_HIT, NECROSIS_HIT])
    assert "- tumor = u\n- necrosis = " in user
    assert "[1] EN: Normal <g1>functioning</g1> neurons\nVI: " + NEURON_VI + "\n[2] EN: Tumor necrosis\nVI: draft two" \
        in user
    assert user.rstrip().endswith("[1] to [2].")
    assert "Glossary" not in build_editor_user(lines, [])


def test_group_hits_keep_only_terms_of_the_group_sources():
    hits = [TUMOR_HIT, NECROSIS_HIT, ("lymph node", "h\u1ea1ch")]
    assert group_hits(["Two tumours", "no other term"], hits) == [TUMOR_HIT]
    assert group_hits(["Lymph-node necrosis"], hits) == [NECROSIS_HIT, ("lymph node", "h\u1ea1ch")]


def test_parse_answer_reads_numbered_lines():
    assert parse_answer("[1] m\u1ed9t\n[2]  hai  \n", 2) == ({1: "m\u1ed9t", 2: "hai"}, [])  # stripped: the guard restores the draft edges
    assert parse_answer("Here are the lines:\n[2] hai\n\n[1] m\u1ed9t", 2) == ({1: "m\u1ed9t", 2: "hai"}, [])


def test_parse_answer_drops_echoes_labels_think_and_fences():
    raw = "<think>x</think>```\n[1] EN: Normal neurons\nVI: m\u1ed9t\n[2] VI: hai\n```"
    assert parse_answer(raw, 2) == ({1: "m\u1ed9t", 2: "hai"}, [])


def test_parse_answer_reports_missing_duplicated_and_out_of_range_lines():
    assert parse_answer("[1] m\u1ed9t\n[3] ba\n[4] b\u1ed1n", 3) == ({1: "m\u1ed9t", 3: "ba"}, [2])
    assert parse_answer("[1] m\u1ed9t\n[1] l\u1ea1i m\u1ed9t\n[2] hai", 2) == ({2: "hai"}, [1])
    assert parse_answer("", 1) == ({}, [1])


def test_parse_answer_joins_a_wrapped_line():
    assert parse_answer("[1] m\u1ed9t d\u00f2ng\nr\u1ea5t d\u00e0i\n[2] hai", 2) == ({1: "m\u1ed9t d\u00f2ng r\u1ea5t d\u00e0i", 2: "hai"}, [])


A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
GUARD_P = ('<a:p %s><a:r><a:rPr lang="en-US"/><a:t>Grade 2 </a:t></a:r><a:r><a:rPr lang="en-US" b="1"/>'
           '<a:t>tumors</a:t></a:r><a:r><a:rPr lang="en-US"/><a:t> with ALK loss in 5 mm</a:t></a:r></a:p>') % A_NS
DRAFT = "C\u00e1c <g1>kh\u1ed1i u</g1> \u0111\u1ed9 2 v\u1edbi m\u1ea5t ALK trong 5 mm"
EDIT = "<g1>Kh\u1ed1i u</g1> \u0111\u1ed9 2 v\u1edbi m\u1ea5t ALK trong 5 mm"


def guard_seg():
    from lxml import etree
    from vitranslate.ooxml.codec import encode
    return encode(etree.fromstring(GUARD_P))


def test_guard_shape_of_the_segment():
    assert guard_seg().tagged == "Grade 2 <g1>tumors</g1> with ALK loss in 5 mm"


def test_guard_accepts_a_clean_edit_and_an_unchanged_line():
    assert guard_line(guard_seg(), DRAFT, EDIT, [], []) == (EDIT, None, [])
    assert guard_line(guard_seg(), DRAFT, DRAFT, [], []) == (DRAFT, None, [])


def test_guard_repairs_tag_spelling_and_restores_the_draft_edges():
    assert guard_line(guard_seg(), DRAFT, EDIT.replace("<g1>", "< g1 >"), [], [])[:2] == (EDIT, None)
    assert guard_line(guard_seg(), " " + DRAFT, EDIT, [], [])[:2] == (" " + EDIT, None)


def test_guard_keeps_the_draft_on_each_failure():
    cases = {
        "": "empty",
        "Kh\u1ed1i u \u0111\u1ed9 2 m\u1ea5t ALK trong 5 mm": "tags",
        "<g1>Kh\u1ed1i u</g1> \u0111\u1ed9 3 m\u1ea5t ALK trong 5 mm": "numbers",
        "<g1>Kh\u1ed1i u</g1> \u0111\u1ed9 2 m\u1ea5t gen trong 5 mm": "acronym",
        "<g1>Kh\u1ed1i u</g1> \u0111\u1ed9 2 m\u1ea5t ALK trong 5 cm": "units",
        "<g1>Kh\u1ed1i u</g1> \u0111\u1ed9 2 m\u1ea5t ALK trong 5 mm, " + "r\u1ea5t " * 20 + "hi\u1ebfm": "ratio",
        "<g1>Kh\u1ed1i u</g1> \u0111\u1ed9 2 m\u1ea5t ALK trong 5 mm \u80bf\u7624": "script",
        "<g1>Kh\u1ed1i u</g1> \u0111\u1ed9 2 m\u1ea5t ALK trong 5 mm [2]": "scaffold",
        "<g1>Kh\u1ed1i u</g1> \u0111\u1ed9 2 VI: m\u1ea5t ALK trong 5 mm": "scaffold",
        "Grade 2 <g1>tumors</g1> with ALK loss in 5 mm": "english",
        "<g1>Kh\u1ed1i u</g1> \u0111\u1ed9 2 m\u1ea5t ALK trong 5 mm pulm\u00f4n": "suspicious",
        "<g1>Kh\u1ed1i u</g1> \u0111\u1ed9 2 m\u1ea5t ALK $\\rightarrow$ trong 5 mm": "latex",
    }
    for edit, reason in cases.items():
        assert guard_line(guard_seg(), DRAFT, edit, [], []) == (DRAFT, reason, []), edit


def test_guard_keeps_canonical_glossary_terms_of_the_draft():
    hits = [("tumor", "kh\u1ed1i u")]
    assert guard_line(guard_seg(), DRAFT, "<g1>U</g1> \u0111\u1ed9 2 m\u1ea5t ALK trong 5 mm", hits, [])[1] == "glossary"
    assert guard_line(guard_seg(), DRAFT, EDIT, hits, [])[1] is None  # case does not matter
    other = DRAFT.replace("kh\u1ed1i u", "b\u01b0\u1edbu")  # term not in the draft: nothing to keep
    assert guard_line(guard_seg(), other, "<g1>U</g1> \u0111\u1ed9 2 m\u1ea5t ALK trong 5 mm", hits, [])[1] is None


def test_guard_short_lines_get_a_few_letters_of_slack():
    seg = guard_seg()
    draft = "C\u00e1c <g1>u</g1> \u0111\u1ed9 2 c\u00f3 ALK, 5 mm"
    assert guard_line(seg, draft, "<g1>U</g1> \u0111\u1ed9 2 c\u00f3 ALK, 5 mm", [], [])[1] is None


def test_guard_reapplies_the_term_lock():
    from vitranslate.packs import TermFix
    fixes = [TermFix("tumor", "kh\u1ed1i u", ["b\u01b0\u1edbu"])]
    edit = "<g1>B\u01b0\u1edbu</g1> \u0111\u1ed9 2 v\u1edbi m\u1ea5t ALK trong 5 mm"
    assert guard_line(guard_seg(), DRAFT, edit, [], fixes) == (EDIT, None, ["tumor"])


def one_run_seg(text):
    from lxml import etree
    from vitranslate.ooxml.codec import encode
    return encode(etree.fromstring('<a:p %s><a:r><a:rPr lang="en-US"/><a:t>%s</a:t></a:r></a:p>' % (A_NS, text)))


def test_guard_reapplies_the_title_order():
    seg = one_run_seg("Slides reviewed by Jane Doe, MD, PhD, in 2024.")
    draft = "Ti\u00eau b\u1ea3n do TS. BS. Jane Doe xem l\u1ea1i n\u0103m 2024."
    edit = "C\u00e1c ti\u00eau b\u1ea3n do BS. TS. Jane Doe xem l\u1ea1i n\u0103m 2024."
    assert guard_line(seg, draft, edit, [], []) == (edit.replace("BS. TS.", "TS. BS."), None, ["title-order"])
    assert guard_line(seg, draft, draft.replace("TS. BS.", "BS TS"), [], []) == (draft, None, [])


def test_guard_accepts_the_house_forms_of_the_active_packs():
    seg = one_run_seg("Many lymphocytes in the node.")
    draft = "Nhi\u1ec1u t\u1ebf b\u00e0o lympho trong h\u1ea1ch."
    edit = "Nhi\u1ec1u lymph\u00f4 b\u00e0o trong h\u1ea1ch."
    assert guard_line(seg, draft, edit, [], []) == (draft, "suspicious", [])
    assert guard_line(seg, draft, edit, [], [], allowed=frozenset({"lymph\u00f4"})) == (edit, None, [])


def editor(replies=None, fn=None, tm=None, fixes=(), hits=()):
    from vitranslate.llm import FakeLLM
    from vitranslate.polish import Polisher
    llm = FakeLLM(replies, fn=fn)
    return Polisher(llm, "EDITOR SYS", "m", "p@1", fixes=fixes, hits=hits, tm=tm), llm


def two_lines():
    seg = guard_seg()
    return [(seg, DRAFT), (seg, DRAFT.replace("\u0111\u1ed9 2", "c\u1ea5p 2"))]


def test_polisher_edits_a_group_in_one_call():
    p, llm = editor(["[1] " + EDIT + "\n[2] " + DRAFT.replace("\u0111\u1ed9 2", "c\u1ea5p 2")])
    results, cached = p.edit(two_lines())
    assert results == [(EDIT, None, []), (DRAFT.replace("\u0111\u1ed9 2", "c\u1ea5p 2"), None, [])] and cached is False
    (system, user), = llm.calls
    assert system == "EDITOR SYS" and "[1] EN: Grade 2 <g1>tumors</g1>" in user and "[2] EN:" in user


def test_polisher_retries_once_when_numbering_is_incomplete():
    p, llm = editor(["[1] " + EDIT, "[1] " + EDIT + "\n[2] " + EDIT])
    results, _ = p.edit(two_lines())
    assert results == [(EDIT, None, []), (EDIT, None, [])]
    assert len(llm.calls) == 2 and "[2]" in llm.calls[1][1].rsplit("\n\n", 1)[1]


def test_polisher_keeps_the_drafts_when_numbering_fails_twice_or_the_model_fails():
    from vitranslate.llm import LLMError
    p, llm = editor(["[1] " + EDIT, "[2] " + EDIT])
    assert p.edit(two_lines())[0] == [(DRAFT, "numbering", []), (DRAFT.replace("\u0111\u1ed9 2", "c\u1ea5p 2"), "numbering", [])]
    assert len(llm.calls) == 2

    def down(system, user):
        raise LLMError("connection refused")

    p, _ = editor(fn=down)
    assert [r[1] for r in p.edit(two_lines())[0]] == ["error", "error"]


def test_polisher_caches_complete_answers_only(tmp_path):
    from vitranslate.tm import TM
    tm = TM(tmp_path / "tm.sqlite")
    answer = "[1] " + EDIT + "\n[2] " + EDIT
    p, _ = editor([answer], tm=tm)
    first = p.edit(two_lines())
    p, llm = editor([], tm=tm)
    assert p.edit(two_lines()) == (first[0], True) and llm.calls == []
    p, llm = editor(["no numbers", "still none"], tm=tm)
    p.edit([(guard_seg(), DRAFT + " ")])  # another group: incomplete twice, nothing stored
    p, llm = editor(["[1] " + EDIT], tm=tm)
    assert p.edit([(guard_seg(), DRAFT + " ")])[1] is False and len(llm.calls) == 1


def fake_editor(system, user):
    """Edits every VI line ("vi ..." -> "vn ..."), except a line whose source mentions "Marker": it gets a broken tag."""
    out, n = [], 0
    for line in user.splitlines():
        if line.startswith("[") and "] EN: " in line:
            n = int(line[1:line.index("]")])
            src = line.split("] EN: ", 1)[1]
        elif line.startswith("VI: "):
            draft = line[4:]
            out.append("[%d] %s" % (n, "vn " + draft[3:] if "Marker" not in src else "vn broken <g9"))
    return "\n".join(out)


def test_run_polish_edits_outcomes_in_place_and_counts(tmp_path):
    from vitranslate.llm import FakeLLM
    from vitranslate.polish import Polisher, run_polish
    items = collect_items(OoxmlPackage(deck(tmp_path / "d.pptx")))
    pairs = drafts(items, {0: ["retried:gate", "glossary_miss:tumor"]})
    first = pairs[0]
    pairs.append((first[0], Outcome(first[1].text, first[1].segment, 0, [])))  # same line again, later part
    llm = FakeLLM(fn=fake_editor)
    events = []
    stats, touched = run_polish(pairs, "pptx", Polisher(llm, "SYS", "m", "k"), events.append,
                                lambda seg, text: ["suspicious:abc"] if text.startswith("Vn Normal") else [])
    assert pairs[0][1].text == "Vn Normal functioning neurons"  # first letter capital like the source
    assert pairs[0][1].flags == ["retried:gate", "suspicious:abc"]  # check flags recomputed for the new text
    assert pairs[3][1].text == "vi Marker panel" and pairs[3][1].flags == ["polish_rejected:tags"]
    assert pairs[6][1].text == "Vn Normal functioning neurons"  # a repeated line reuses the first decision
    assert len(llm.calls) == 3 and sorted(touched) == [0, 1, 2, 3, 4, 5, 6]
    assert stats == {"enabled": True, "groups": 4, "calls": 3, "cached": 0, "lines": 6, "changed": 5, "reused": 1,
                     "rejected": {"tags": 1}}
    assert events[0] == {"stage": "polish", "done": 0, "total": 4}
    assert events[-1] == {"stage": "polish", "done": 4, "total": 4}


def plain_seg(text):
    from lxml import etree
    from vitranslate.ooxml.codec import encode
    return encode(etree.fromstring('<a:p %s><a:r><a:rPr lang="en-US"/><a:t>%s</a:t></a:r></a:p>' % (A_NS, text)))


def test_guard_keeps_the_final_period_of_the_draft_unless_the_source_agrees():
    seg = plain_seg("Alzheimer disease remains the leading cause of dementia")
    draft = "B\u1ec7nh Alzheimer l\u00e0 nguy\u00ean nh\u00e2n ph\u1ed5 bi\u1ebfn nh\u1ea5t g\u00e2y ra sa s\u00fat tr\u00ed tu\u1ec7"
    edit = "B\u1ec7nh Alzheimer l\u00e0 nguy\u00ean nh\u00e2n ph\u1ed5 bi\u1ebfn nh\u1ea5t g\u00e2y sa s\u00fat tr\u00ed tu\u1ec7"
    assert guard_line(seg, draft, edit + ".", [], [])[:2] == (edit, None)
    seg = plain_seg("Grade is common.")
    assert guard_line(seg, "\u0110\u1ed9 th\u01b0\u1eddng g\u1eb7p.", "Ph\u00e2n \u0111\u1ed9 th\u01b0\u1eddng g\u1eb7p", [], [])[:2] == ("Ph\u00e2n \u0111\u1ed9 th\u01b0\u1eddng g\u1eb7p.", None)
    seg = plain_seg("Grade is common")
    assert guard_line(seg, "\u0110\u1ed9 th\u01b0\u1eddng g\u1eb7p.", "Ph\u00e2n \u0111\u1ed9 th\u01b0\u1eddng g\u1eb7p", [], [])[:2] == ("Ph\u00e2n \u0111\u1ed9 th\u01b0\u1eddng g\u1eb7p", None)


def test_guard_first_letter_case_flip_follows_the_source():
    assert guard_line(plain_seg("Normal"), "B\u00ecnh th\u01b0\u1eddng", "b\u00ecnh th\u01b0\u1eddng", [], [])[:2] == ("B\u00ecnh th\u01b0\u1eddng", None)
    assert guard_line(plain_seg("other"), "kh\u00e1c", "Kh\u00e1c", [], [])[:2] == ("kh\u00e1c", None)
    assert guard_line(plain_seg("The cingulate gyrus"), "h\u1ed3i \u0111ai", "H\u1ed3i \u0111ai", [], [])[:2] == ("H\u1ed3i \u0111ai", None)
    # the first word changed: the case is the editor's, but a source in capitals starts the line with a capital
    assert guard_line(plain_seg("other areas"), "c\u00e1c v\u00f9ng kh\u00e1c", "V\u00f9ng kh\u00e1c", [], [])[:2] == ("V\u00f9ng kh\u00e1c", None)
    assert guard_line(plain_seg("Other features"), "C\u00e1c \u0111\u1eb7c \u0111i\u1ec3m kh\u00e1c",
                      "\u0111\u1eb7c \u0111i\u1ec3m kh\u00e1c", [], []) == ("\u0110\u1eb7c \u0111i\u1ec3m kh\u00e1c", None, [])


def test_guard_rejects_an_edit_that_only_drops_words_carrying_meaning():
    seg = plain_seg("Autopsy brain only")
    assert guard_line(seg, "Ch\u1ec9 t\u1eed thi\u1ebft n\u00e3o", "N\u00e3o t\u1eed thi\u1ebft", [], [])[1] == "dropped"
    seg = plain_seg("Insoluble aggregates of")
    assert guard_line(seg, "C\u00e1c k\u1ebft t\u1eadp kh\u00f4ng tan c\u1ee7a", "K\u1ebft t\u1eadp kh\u00f4ng tan c\u1ee7a", [], [])[1] is None
    seg = plain_seg("Inversely related to the severity of AD pathology.")
    draft = "T\u1ef7 l\u1ec7 ngh\u1ecbch v\u1edbi m\u1ee9c \u0111\u1ed9 c\u1ee7a b\u1ec7nh l\u00fd Alzheimer (AD)."
    assert guard_line(seg, draft, draft.replace("c\u1ee7a ", ""), [], [])[1] is None
    seg = plain_seg("Consider a modified approach in a different setting")
    draft = "Xem x\u00e9t m\u1ed9t c\u00e1ch ti\u1ebfp c\u1eadn s\u1eeda \u0111\u1ed5i trong b\u1ed1i c\u1ea3nh kh\u00e1c"
    assert guard_line(seg, draft, "Xem x\u00e9t c\u00e1ch ti\u1ebfp c\u1eadn s\u1eeda \u0111\u1ed5i trong b\u1ed1i c\u1ea3nh", [], [])[1] == "dropped"
    # a reworded line (new words) is not a pure deletion
    seg = plain_seg("Fluctuation is common")
    assert guard_line(seg, "Bi\u1ebfn \u0111\u1ed9ng l\u00e0 ph\u1ed5 bi\u1ebfn", "Bi\u1ebfn \u0111\u1ed9ng th\u01b0\u1eddng g\u1eb7p", [], [])[1] is None


def test_guard_lets_the_editor_remove_a_repeated_word():
    seg = plain_seg("TERT promoter mutations appear in progression")
    draft = ("\u0110\u1ed9t bi\u1ebfn v\xf9ng kh\u1edfi \u0111\u1ed9ng TERT TERT xu\u1ea5t hi\u1ec7n khi b\u1ec7nh "
             "ti\u1ebfn tri\u1ec3n")
    assert guard_line(seg, draft, draft.replace("TERT TERT", "TERT"), [], [])[1] is None


def test_guard_keeps_a_glossary_term_written_without_its_english_note():
    hits = [("clear cell", "t\u1ebf b\u00e0o s\u00e1ng (clear cell)")]
    seg = plain_seg("Clear cell renal cell carcinoma (CCRCC)")
    draft = "Ung th\u01b0 bi\u1ec3u m\u00f4 t\u1ebf b\u00e0o s\u00e1ng th\u1eadn (CCRCC)"
    assert guard_line(seg, draft, "Ung th\u01b0 bi\u1ec3u m\u00f4 t\u1ebf b\u00e0o th\u1eadn d\u1ea1ng s\u00e1ng (CCRCC)", hits, [])[1] == "glossary"


def test_guard_rejects_english_words_copied_back_from_the_source():
    seg = plain_seg("Adrenal cortical carcinoma with sarcomatoid features")
    draft = "Ung th\u01b0 bi\u1ec3u m\xf4 v\u1ecf th\u01b0\u1ee3ng th\u1eadn k\xe8m th\xe0nh ph\u1ea7n d\u1ea1ng sarc\xf4m"
    edit = "Ung th\u01b0 bi\u1ec3u m\xf4 v\u1ecf th\u01b0\u1ee3ng th\u1eadn k\xe8m th\xe0nh ph\u1ea7n d\u1ea1ng sarcomatoid"
    assert guard_line(seg, draft, edit, [], [])[1] == "english"
    # an English note that the glossary itself writes is fine
    seg = plain_seg("TERT promoter mutation")
    hits = [("TERT promoter mutation", "\u0111\u1ed9t bi\u1ebfn v\u00f9ng kh\u1edfi \u0111\u1ed9ng (promoter) TERT")]
    assert guard_line(seg, "\u0110\u1ed9t bi\u1ebfn TERT v\u00f9ng kh\u1edfi \u0111\u1ed9ng", "\u0110\u1ed9t bi\u1ebfn v\u00f9ng kh\u1edfi \u0111\u1ed9ng (promoter) TERT", hits,
                      [])[1] is None


def test_guard_a_pure_case_change_keeps_the_draft_casing_apart_from_the_first_letter():
    assert guard_line(plain_seg("Prognosis"), "Ti\u00ean l\u01b0\u1ee3ng", "TI\u00caN L\u01af\u1ee2NG", [], [])[:2] == ("Ti\u00ean l\u01b0\u1ee3ng", None)
    assert guard_line(plain_seg("Cerebellar atrophy"), "teo ti\u1ec3u n\u00e3o", "TEO TI\u1ec2U N\u00c3O", [], [])[:2] == (
        "Teo ti\u1ec3u n\u00e3o", None)


def test_guard_free_words_cover_active_voice_and_trong_viec():
    seg = plain_seg("Difficulties in assessing margins")
    draft = "Nh\u1eefng kh\u00f3 kh\u0103n trong vi\u1ec7c \u0111\u00e1nh gi\u00e1 di\u1ec7n c\u1eaft"
    assert guard_line(seg, draft, "Kh\u00f3 kh\u0103n \u0111\u00e1nh gi\u00e1 di\u1ec7n c\u1eaft", [], [])[1] is None
    seg = plain_seg("Tumours located in the lung")
    assert guard_line(seg, "U n\u1eb1m trong ph\u1ed5i", "U n\u1eb1m ph\u1ed5i", [], [])[1] == "dropped"  # "trong" alone carries meaning
    seg = plain_seg("Size was measured by two pathologists.")
    draft = "K\u00edch th\u01b0\u1edbc \u0111\u00e3 \u0111\u01b0\u1ee3c \u0111o b\u1edfi hai b\u00e1c s\u0129."
    assert guard_line(seg, draft, "K\u00edch th\u01b0\u1edbc \u0111o hai b\u00e1c s\u0129.", [], [])[1] is None


LYMPH_SRC = "Small clusters of mature lymphocytes surround the vessel wall in the deep dermis of the biopsy"
LYMPH_VI = ("C\u00e1c c\u1ee5m nh\u1ecf t\u1ebf b\u00e0o lympho tr\u01b0\u1edfng th\u00e0nh bao quanh th\u00e0nh "
            "m\u1ea1ch \u1edf l\u1edbp b\u00ec s\u00e2u c\u1ee7a m\u1eabu sinh thi\u1ebft")
LYMPH_PHRASE = "t\u1ebf b\u00e0o lympho tr\u01b0\u1edfng th\u00e0nh bao quanh"  # 7 syllables


def test_guard_rejects_a_phrase_the_edit_says_twice():
    seg = plain_seg(LYMPH_SRC)
    assert guard_line(seg, LYMPH_VI, LYMPH_VI[4:] + " " + LYMPH_PHRASE, [], [])[1] == "repeat"
    seg = plain_seg("Nodule in the left lobe of the gland")
    draft = "N\u1ed1t \u1edf th\u00f9y tr\u00e1i c\u1ee7a tuy\u1ebfn"
    stutter = "N\u1ed1t \u1edf th\u00f9y tr\u00e1i th\u00f9y tr\u00e1i c\u1ee7a tuy\u1ebfn"
    assert guard_line(seg, draft, stutter, [], [])[1] == "repeat"


def test_guard_allows_repeats_the_draft_had_and_reused_short_terms():
    seg = plain_seg("Clinical history, clinical history!")
    draft = "B\u1ec7nh s\u1eed l\u00e2m s\u00e0ng, b\u1ec7nh s\u1eed l\u00e2m s\u00e0ng!"
    edit = "Ti\u1ec1n s\u1eed l\u00e2m s\u00e0ng, ti\u1ec1n s\u1eed l\u00e2m s\u00e0ng!"
    assert guard_line(seg, draft, edit, [], [])[1] is None
    seg = plain_seg("Both upper and lower pole lesions")
    draft = "T\u1ed5n th\u01b0\u01a1ng c\u1ef1c tr\u00ean v\u00e0 c\u1ef1c d\u01b0\u1edbi"
    edit = "T\u1ed5n th\u01b0\u01a1ng c\u1ef1c tr\u00ean v\u00e0 t\u1ed5n th\u01b0\u01a1ng c\u1ef1c d\u01b0\u1edbi"
    assert guard_line(seg, draft, edit, [], [])[1] is None


def test_guard_keeps_footnote_and_reference_symbols_of_the_source():
    data = "D\u1eef li\u1ec7u t\u1eeb m\u1ed9t trung t\u00e2m"
    assert guard_line(plain_seg("\u2020Data from a single centre"), "\u2020" + data, data, [], [])[1] == "symbols"
    assert guard_line(plain_seg("Seen in 12 cases\u00b2"), "G\u1eb7p \u1edf 12 ca\u00b2", "G\u1eb7p \u1edf 12 ca", [],
                      [])[1] == "symbols"
    reported = "\u0110\u00e3 \u0111\u01b0\u1ee3c b\u00e1o c\u00e1o tr\u01b0\u1edbc \u0111\u00e2y"
    assert guard_line(plain_seg("Reported before [3]"), reported + " [3]", reported + " (3)", [], [])[1] == "symbols"
    # a symbol the source does not have may go
    assert guard_line(plain_seg("Data from a single centre"), "*" + data, data, [], [])[1] is None


FTC_DRAFT = ("C\u1eaft th\u00f9y tuy\u1ebfn gi\u00e1p do ung th\u01b0 bi\u1ec3u m\u00f4 d\u1ea1ng nang (FTC)")
FTC_WEAK = "C\u1eaft th\u00f9y tuy\u1ebfn gi\u00e1p do u bi\u1ec3u m\u00f4 d\u1ea1ng nang (FTC)"


def test_guard_rejects_a_malignancy_term_weakened_to_tumour():
    assert guard_line(plain_seg("Lobectomy for FTC"), FTC_DRAFT, FTC_WEAK, [], [])[1] == "term_weakened"
    seg = plain_seg("how invasive the cancer is")
    draft = "m\u1ee9c \u0111\u1ed9 x\u00e2m l\u1ea5n c\u1ee7a ung th\u01b0"
    assert guard_line(seg, draft, "m\u1ee9c \u0111\u1ed9 x\u00e2m l\u1ea5n c\u1ee7a kh\u1ed1i u", [], [])[1] == "term_weakened"


def test_guard_lets_a_tumour_overtranslated_as_cancer_become_tumour():
    seg = plain_seg("Tumours of the salivary gland")
    draft = "Ung th\u01b0 tuy\u1ebfn n\u01b0\u1edbc b\u1ecdt"
    assert guard_line(seg, draft, "U tuy\u1ebfn n\u01b0\u1edbc b\u1ecdt", [], [])[1] is None
    seg = plain_seg("A tumour next to a carcinoma")
    draft = "Ung th\u01b0 c\u1ea1nh ung th\u01b0 bi\u1ec3u m\u00f4"
    assert guard_line(seg, draft, "U c\u1ea1nh ung th\u01b0 bi\u1ec3u m\u00f4", [], [])[1] is None


def test_guard_rejects_an_anatomical_site_the_source_does_not_name():
    seg = plain_seg("CT: a 3 cm renal mass on the right")
    draft = "CT: kh\u1ed1i th\u1eadn 3 cm b\u00ean ph\u1ea3i"
    assert guard_line(seg, draft, "CT: kh\u1ed1i b\u1ec3 th\u1eadn 3 cm b\u00ean ph\u1ea3i", [], [])[1] == "anatomy"


def test_guard_accepts_a_site_the_source_names_even_as_a_word_stem():
    seg = plain_seg("Nodule in the left kidney")
    assert guard_line(seg, "N\u1ed1t \u1edf b\u00ean tr\u00e1i", "N\u1ed1t \u1edf th\u1eadn tr\u00e1i", [], [])[1] is None
    seg = plain_seg("Status after nephrectomy")
    draft = "T\u00ecnh tr\u1ea1ng sau ph\u1eabu thu\u1eadt"
    assert guard_line(seg, draft, "T\u00ecnh tr\u1ea1ng sau c\u1eaft th\u1eadn", [], [])[1] is None


FNA_HIT = ("fine needle aspiration", "ch\u1ecdc h\u00fat kim nh\u1ecf")


def test_glossary_rows_count_for_their_acronym_in_the_source():
    seg = plain_seg("FNA RESULTS BY YEAR")
    draft = "K\u1ebft qu\u1ea3 ch\u1ecdc h\u00fat kim nh\u1ecf theo n\u0103m"
    edit = "K\u1ebft qu\u1ea3 sinh thi\u1ebft kim theo n\u0103m"
    assert group_hits(["FNA RESULTS BY YEAR"], [FNA_HIT, TUMOR_HIT]) == [FNA_HIT]
    assert guard_line(seg, draft, edit, [FNA_HIT], [])[1] == "glossary"
    assert group_hits(["Fna results"], [FNA_HIT]) == []  # the acronym is an upper-case word of its own


def test_guard_short_lines_may_shrink_more():
    seg = plain_seg("Presence of necrosis")
    assert guard_line(seg, "S\u1ef1 hi\u1ec7n di\u1ec7n c\u1ee7a ho\u1ea1i t\u1eed", "C\u00f3 ho\u1ea1i t\u1eed", [], [])[1] is None
    seg = plain_seg("Lymph node metastases were present in eight of the twenty patients studied")
    draft = "Di c\u0103n h\u1ea1ch b\u1ea1ch huy\u1ebft hi\u1ec7n di\u1ec7n \u1edf t\u00e1m trong s\u1ed1 hai m\u01b0\u01a1i b\u1ec7nh nh\u00e2n \u0111\u01b0\u1ee3c nghi\u00ean c\u1ee9u"
    assert guard_line(seg, draft, "C\u00f3 di c\u0103n h\u1ea1ch \u1edf t\u00e1m b\u1ec7nh nh\u00e2n", [], [])[1] == "ratio"
