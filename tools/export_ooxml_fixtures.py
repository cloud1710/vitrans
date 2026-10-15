"""Dump what the Python OOXML layer does with DOCX / PPTX files, as JSON, for the JavaScript port's parity tests.

For each file: the selected parts and labels, the item list of collect_items (part, index, location, skip reason, own
text and a summary of the segment: tagged text, dominant run, spans, fixed elements, anchors, locked text) and, for
three deterministic write-backs, what apply() + fit_pptx() leave in the package:
  identity  the tagged source applied back unchanged
  fake      every word of the tagged source prefixed "VI_" (tags untouched)
  messy     the fake text with placeholders moved / duplicated / dropped / unknown, an unknown span, control
            characters, CR LF, NFD text and a lone surrogate (exercises the apply() fixes and the text cleaning)
Per scenario: the text applied to each segment, the fixes apply() reported, the fit changes and, for every changed
part, the sha256 of the bytes the helper would write plus a canonical form of that XML (see canon_xml; the JavaScript
side computes the same form from its own output).

Usage (Python = the vitranslate env):
  export_ooxml_fixtures.py fixtures            build the synthetic files into web/engine/tests/fixtures and dump them
  export_ooxml_fixtures.py dump OUT FILE...    dump the given files into OUT (one <stem>.json each)
  export_ooxml_fixtures.py corpus OUT          dump the files of private/corpus.txt into OUT as corpus_<n>.json
                                               (LOCAL ONLY: OUT must stay outside the repository)
Console output is ASCII only.
"""
import base64
import hashlib
import io
import json
import re
import struct
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "helper"))
sys.path.insert(0, str(ROOT / "helper" / "tests"))

from lxml import etree  # noqa: E402

from vitranslate.fit import fit_pptx  # noqa: E402
from vitranslate.ooxml.codec import apply, encode  # noqa: E402
from vitranslate.ooxml.document import collect_items  # noqa: E402
from vitranslate.ooxml.ns import A, P, R, W  # noqa: E402
from vitranslate.ooxml.package import OoxmlPackage  # noqa: E402
from vitranslate.ooxml.parts import part_labels, select_parts  # noqa: E402
from vitranslate.tags import TOKEN_RE  # noqa: E402

FIXTURES = ROOT / "web" / "engine" / "tests" / "fixtures"
SCENARIOS = ("identity", "fake", "messy")
_PARSER = etree.XMLParser(remove_blank_text=False, resolve_entities=False)


# ---------------------------------------------------------------- canonical forms (mirrored in tests/canon.mjs)

def esc(s):
    """Quoted string with every character outside printable ASCII (and the quote / backslash) as \\uXXXX of its
    UTF-16 code units, so Python and JavaScript produce the same text."""
    out = ['"']
    for ch in s:
        o = ord(ch)
        if 0x20 <= o < 0x7F and ch not in '"\\':
            out.append(ch)
        elif o > 0xFFFF:
            o -= 0x10000
            out.append("\\u%04x\\u%04x" % (0xD800 + (o >> 10), 0xDC00 + (o & 0x3FF)))
        else:
            out.append("\\u%04x" % o)
    out.append('"')
    return "".join(out)


def _ns_local(key):
    q = etree.QName(key)
    return q.namespace or "", q.localname


def _attrs(el):
    return sorted(_ns_local(k) + (v,) for k, v in el.attrib.items())


def _qname(el):
    return (el.prefix + ":" if el.prefix else "") + etree.QName(el).localname


def canon_xml(data):
    """Canonical lines of an XML document: top-level comments / PIs, elements with their prefix and namespace, the
    namespace declarations that change the scope, attributes sorted by (namespace, local name), text, comments and
    PIs in order. Redundant declarations, attribute order, quoting and escaping do not show."""
    root = etree.fromstring(data, parser=_PARSER)
    lines = []
    before, n = [], root.getprevious()
    while n is not None:
        before.insert(0, n)
        n = n.getprevious()
    after, n = [], root.getnext()
    while n is not None:
        after.append(n)
        n = n.getnext()

    def other(node):
        if isinstance(node, etree._Comment):
            lines.append("C " + esc(node.text or ""))
        elif isinstance(node, etree._ProcessingInstruction):
            lines.append("P " + node.target + " " + esc(node.text or ""))
        else:
            lines.append("? " + esc(str(node)))

    def walk(el, scope):
        ns, loc = _ns_local(el.tag)
        lines.append("E " + _qname(el) + " " + esc(ns))
        own = {}
        for prefix, uri in el.nsmap.items():
            if scope.get(prefix or "") != uri:
                own[prefix or ""] = uri
        for prefix in sorted(own):
            lines.append("N " + prefix + "=" + esc(own[prefix]))
        inner = dict(scope)
        inner.update(own)
        for ans, aloc, v in _attrs(el):
            lines.append("A " + esc(ans) + " " + aloc + "=" + esc(v))
        if el.text:
            lines.append("T " + esc(el.text))
        for ch in el:
            if isinstance(ch.tag, str):
                walk(ch, inner)
            else:
                other(ch)
            if ch.tail:
                lines.append("T " + esc(ch.tail))
        lines.append("/E")

    for node in before:
        other(node)
    walk(root, {})
    for node in after:
        other(node)
    return "\n".join(lines)


def ecanon(el):
    """Compact form of a (detached) element subtree, without namespace declarations or its own tail."""
    if el is None:
        return None

    def walk(e):
        if not isinstance(e.tag, str):
            if isinstance(e, etree._Comment):
                return "<!--" + esc(e.text or "") + "-->"
            return "<?" + e.target + " " + esc(e.text or "") + "?>"
        ns, loc = _ns_local(e.tag)
        s = "<{" + ns + "}" + loc + "".join(" {%s}%s=%s" % (a, b, esc(v)) for a, b, v in _attrs(e)) + ">"
        if e.text:
            s += esc(e.text)
        for ch in e:
            s += walk(ch)
            if ch.tail:
                s += esc(ch.tail)
        return s + "</>"
    return walk(el)


def desc(node):
    if isinstance(node.tag, str):
        ns, loc = _ns_local(node.tag)
        return "{" + ns + "}" + loc
    if isinstance(node, etree._Comment):
        return "#comment"
    return "#pi " + node.target


def ckey(el):
    if el is None:
        return None
    ns, loc = _ns_local(el.tag)
    return "{" + ns + "}" + loc + "".join(" {%s}%s=%s" % (a, b, esc(v)) for a, b, v in _attrs(el))


def span_summary(span):
    return {"sig": span.sig, "rpr": ecanon(span.rpr), "container": ckey(span.container)}


def seg_summary(seg):
    if seg is None:
        return None
    return {
        "kind": seg.kind, "tagged": seg.tagged, "locked": list(seg.locked),
        "dominant": span_summary(seg.dominant),
        "spans": {str(n): span_summary(s) for n, s in seg.spans.items()},
        "fixed": {str(n): {"wrap": f.wrap, "nodes": [desc(e) for e in f.elements], "container": ckey(f.container),
                           "wrap_rpr": ecanon(f.wrap_rpr)} for n, f in seg.fixed.items()},
        "anchors_start": [desc(a) for a in seg.anchors_start],
        "anchors_end": [desc(a) for a in seg.anchors_end],
    }


# ---------------------------------------------------------------- deterministic write-backs

def fake_text(tagged):
    """Every whitespace-separated word of the plain pieces gets the prefix VI_; tags stay where they are."""
    out, pos = [], 0
    for m in TOKEN_RE.finditer(tagged):
        out.append(re.sub(r"(\S+)", r"VI_\1", tagged[pos:m.start()]))
        out.append(m.group(0))
        pos = m.end()
    out.append(re.sub(r"(\S+)", r"VI_\1", tagged[pos:]))
    return "".join(out)


def messy_text(tagged, n):
    text = fake_text(tagged)
    xs = re.findall(r"<x\d+/>", text)
    if xs:
        first = xs[0]
        text = text.replace(first, "", 1)
        text = first + text                    # first placeholder moved to the front
        text = text + xs[-1]                   # last placeholder duplicated at the end
        if len(xs) > 2:
            text = text.replace(xs[1], "", 1)  # second placeholder dropped (apply appends it)
    if n % 2 == 0:
        text += " <x99/>"                      # unknown placeholder
    gs = re.findall(r"<g(\d+)>", text)
    if gs and n % 3 == 0:
        text = text.replace("<g%s>" % gs[0], "<g%s><g77>" % gs[0], 1).replace(
            "</g%s>" % gs[0], "</g77></g%s>" % gs[0], 1)  # unknown span nested in a known one
    noise = ["A\r\nB", "ma" + "\u0300" + "ng", "C\vD", "E\x01F", "G" + chr(0xD800) + "H", "\u00a0"]
    return noise[n % len(noise)] + " " + text


def scenario_text(name, tagged, n):
    if name == "identity":
        return tagged
    if name == "fake":
        return fake_text(tagged)
    return messy_text(tagged, n)


def serialize(pkg, name):
    return etree.tostring(pkg._roots[name].getroottree(), xml_declaration=True, encoding="UTF-8", standalone=True)


def dump(path, notes=True, comments=False, source=None):
    pkg = OoxmlPackage(path)
    items = collect_items(pkg, notes=notes, comments=comments)
    parts = select_parts(pkg, notes=notes, comments=comments)
    labels = part_labels(pkg)
    out = {
        "file": source or Path(path).name, "kind": pkg.kind, "options": {"notes": notes, "comments": comments},
        "names": pkg.names(), "parts": parts, "labels": {p: labels[p] for p in parts},
        "items": [{"part": i.part, "index": i.index, "location": i.location, "skip": i.skip, "text": i.text,
                   "seg": seg_summary(i.segment)} for i in items],
        "scenarios": {},
    }
    for name in SCENARIOS:
        pkg = OoxmlPackage(path)
        work = [i for i in collect_items(pkg, notes=notes, comments=comments) if i.segment]
        texts, fixes, pairs = [], [], []
        for n, item in enumerate(work):
            text = scenario_text(name, item.segment.tagged, n)
            texts.append(text)
            fixes.append(apply(item.element, item.segment, text))
            pkg.mark_dirty(item.part)
            pairs.append((item, SimpleNamespace(text=text)))
        fit = fit_pptx([(i, o) for i, o in pairs if "/slides/" in i.part]) if pkg.kind == "pptx" else []
        changed = {}
        for part in pkg.names():
            if part in pkg.dirty and part not in changed:
                data = serialize(pkg, part)
                changed[part] = {"sha256": hashlib.sha256(data).hexdigest(), "canon": canon_xml(data)}
        out["scenarios"][name] = {"texts": texts, "fixes": fixes, "fit": fit, "changed": changed}
    return out


def _xml_cases():
    """Single XML parts: parse + tostring as the helper does (or the error lxml raises)."""
    e_acute, smile = chr(0xE9), chr(0x1F600)
    return {
        "cdata": b'<r><t><![CDATA[a < b & c]]></t></r>',
        "entities": b'<r a="x &amp; &lt; &gt; &quot; \'"><t>&amp;&lt;&gt;&quot;&apos;&#x1F600;&#160;</t></r>',
        "cr_ref_text": b'<r><t>a&#13;b&#13;&#10;c</t></r>',
        "crlf_raw_text": b'<r><t>a\r\nb\rc</t></r>',
        "attr_ws": b'<r a="x&#9;y&#10;z&#13;w" b="lit\ttab\nnl"/>',
        "attr_gt": b'<r a="1 > 0"/>',
        "comment_pi_top": b'<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n<!-- top --><?pi data?><r/>'
                          b'<!-- after -->',
        "pi_no_data": b'<r><?go?></r>',
        "empty_elem": b'<r><a></a><b/></r>',
        "bom": b'\xef\xbb\xbf<?xml version="1.0" encoding="UTF-8"?><r>x</r>',
        "utf16": ('<?xml version="1.0" encoding="UTF-16"?><r>' + e_acute + smile + '</r>').encode("utf-16"),
        "nel_ls": ('<r><t>a' + chr(0x85) + 'b' + chr(0x2028) + 'c</t></r>').encode("utf-8"),
        "xmlns_empty": b'<r xmlns="urn:a"><c xmlns=""/></r>',
        "redundant_ns": ('<w:p xmlns:w="%s"><w:r xmlns:w="%s"><w:t>x</w:t></w:r></w:p>' % (W, W)).encode(),
        "ns_attr_only": b'<r xmlns:q="urn:q"><c q:a="1"/></r>',
        "prefix_rebind": b'<a:r xmlns:a="urn:1"><a:c xmlns:a="urn:2"><a:d/></a:c></a:r>',
        "standalone_no_decl": b'<?xml version="1.0"?><r/>',
        "ws_only_root_text": b'<r>\n  <a/>\n</r>\n',
        "latin1": ('<?xml version="1.0" encoding="ISO-8859-1"?><r>' + e_acute + '</r>').encode("latin-1"),
        # ill-formed: lxml raises XMLSyntaxError, the engine must raise XmlParseError (never write such a part)
        "bad_mismatched_tag": b'<w:p xmlns:w="urn:w"><w:r><w:t>x</w:t></w:p>',
        "bad_truncated": b'<w:document xmlns:w="urn:w"><w:body><w:p><w:r><w:t>Hel',
        "bad_undefined_entity": b'<r><t>a &nbsp; b</t></r>',
        "bad_no_root": b'<?xml version="1.0"?>\n<!-- only a comment -->',
        "bad_empty": b'',
        "bad_two_roots": b'<a/><b/>',
        "bad_attribute_twice": b'<a b="1" b="2"/>',
        "bad_utf8": b'<r>caf\xc3</r>',
    }


_CLONE_LINK = ('<w:hyperlink xmlns:q1="urn:q1" xmlns:q2="urn:q2" q1:a="1" q2:b="2" w:history="1"><w:r><w:rPr><w:b/>'
               '</w:rPr><w:t>bold</w:t></w:r><w:r><w:t xml:space="preserve"> linked words</w:t></w:r></w:hyperlink>')
_CLONE_PARA = ('<w:p%s><w:r><w:t xml:space="preserve">Read the </w:t></w:r>%s<w:r><w:t xml:space="preserve"> now and '
               'then more text</w:t></w:r></w:p>')
_CLONE_TEXT = "Read <g1>bold</g1> mid <g2>linked</g2> now and then more text"


def _clone_cases():
    """Split hyperlinks: the clone takes the hyperlink's attributes; namespaces not in scope get new ns<N> prefixes,
    counted per document as lxml does."""
    w_decl = ' xmlns:w="%s"' % W
    local_r = ('<w:hyperlink xmlns:r="%s" r:id="rId1"><w:r><w:rPr><w:b/></w:rPr><w:t>bold</w:t></w:r><w:r>'
               '<w:t xml:space="preserve"> linked words</w:t></w:r></w:hyperlink>' % R)
    return {
        "two_foreign": {"xml": _CLONE_PARA % (w_decl, _CLONE_LINK), "texts": [_CLONE_TEXT]},
        "local_r": {"xml": _CLONE_PARA % (w_decl, local_r), "texts": [_CLONE_TEXT]},
        "three_pieces": {"xml": _CLONE_PARA % (w_decl, _CLONE_LINK),
                         "texts": ["<g1>bold</g1> a <g2>linked</g2> b <g1>bold again</g1> end"]},
        "two_paragraphs": {"xml": '<w:body%s>%s%s</w:body>' % (w_decl, _CLONE_PARA % ("", _CLONE_LINK),
                                                               _CLONE_PARA % ("", _CLONE_LINK)),
                           "texts": [_CLONE_TEXT, _CLONE_TEXT]},
    }


def lxml_cases():
    """Reference results of lxml / the helper for the engine's unit checks (tests/checks.mjs)."""
    xml = {}
    for name, data in _xml_cases().items():
        case = {"in": base64.b64encode(data).decode("ascii")}
        try:
            root = etree.fromstring(data, parser=_PARSER)
            out = etree.tostring(root.getroottree(), xml_declaration=True, encoding="UTF-8", standalone=True)
            case["out"] = base64.b64encode(out).decode("ascii")
        except etree.XMLSyntaxError:
            case["error"] = "XMLSyntaxError"
        if name.startswith("bad_") != ("error" in case):
            raise AssertionError("lxml outcome not as named: " + name)
        xml[name] = case
    clone = {}
    for name, case in _clone_cases().items():
        root = etree.fromstring(case["xml"])
        paras = [root] if root.tag == "{%s}p" % W else root.findall("{%s}p" % W)
        tagged, fixes = [], []
        for p, text in zip(paras, case["texts"]):
            seg = encode(p)
            tagged.append(seg.tagged)
            fixes.append(apply(p, seg, text))
        clone[name] = dict(case, tagged=tagged, fixes=fixes, py=etree.tostring(root, encoding="unicode"))
    return {"xml": xml, "clone": clone}


def write_json(obj, path):
    Path(path).write_text(json.dumps(obj, ensure_ascii=True, indent=1, sort_keys=False) + "\n", encoding="utf-8")


# ---------------------------------------------------------------- synthetic files (synthetic text only)

def _codec_w_paragraphs():
    """The WordprocessingML paragraphs of helper/tests/test_codec.py (synthetic text)."""
    ns = ('xmlns:w="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
          'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"' % W)
    head = "<w:p %s>" % ns
    return [
        head + '<w:pPr><w:pStyle w:val="Normal"/></w:pPr>\n  <w:bookmarkStart w:id="0" w:name="_GoBack"/>\n'
        '  <w:r w:rsidR="00A1"><w:rPr><w:lang w:val="en-US"/></w:rPr><w:t xml:space="preserve">See </w:t></w:r>\n'
        '  <w:proofErr w:type="spellStart"/>\n'
        '  <w:hyperlink r:id="rIdLINK1"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/></w:rPr><w:t>the WHO site</w:t>'
        '</w:r></w:hyperlink>\n  <w:proofErr w:type="spellEnd"/>\n  <w:r><w:t xml:space="preserve"> for </w:t></w:r>\n'
        '  <w:r><w:rPr><w:b/></w:rPr><w:t>grading</w:t></w:r>\n  <w:r><w:tab/><w:t>criteria</w:t></w:r>\n'
        '  <w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:t>1</w:t></w:r>\n'
        '  <w:r><w:fldChar w:fldCharType="begin"/></w:r>\n'
        '  <w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>\n'
        '  <w:r><w:fldChar w:fldCharType="separate"/></w:r>\n  <w:r><w:t>4</w:t></w:r>\n'
        '  <w:r><w:fldChar w:fldCharType="end"/></w:r>\n  <w:bookmarkEnd w:id="0"/>\n</w:p>',
        head + '<w:r><w:t xml:space="preserve">Grade </w:t></w:r><w:commentRangeStart w:id="3"/>'
        '<w:smartTag w:uri="urn:x" w:element="place"><w:smartTagPr><w:attr w:name="a" w:val="b"/></w:smartTagPr>'
        '<w:r><w:rPr><w:b/></w:rPr><w:t>Hanoi</w:t></w:r></w:smartTag><w:r><w:tab/><w:t>tumour</w:t></w:r>'
        '<w:commentRangeEnd w:id="3"/></w:p>',
        head + '<w:r><w:rPr><w:rFonts w:hint="eastAsia"/><w:b/><w:eastAsianLayout w:id="1"/></w:rPr>'
        '<w:t>Text with layout</w:t></w:r></w:p>',
        head + '<w:r><w:rPr><w:b/><w14:ligatures w14:val="standardContextual"/></w:rPr><w:t>Text ligatures</w:t>'
        '</w:r></w:p>',
        head + '<w:r><w:rPr><w:b/></w:rPr><w:t>a</w:t></w:r><w:r><w:rPr><w:u w:val="single"/></w:rPr>'
        '<w:t xml:space="preserve"> </w:t></w:r><w:r><w:t>long dominant text</w:t></w:r></w:p>',
        head + '<w:r><w:rPr><w:rFonts w:ascii="Arial"/><w:sz w:val="24"/></w:rPr><w:t xml:space="preserve">Tumour '
        '</w:t></w:r><w:r><w:rPr><w:rFonts w:ascii="Times New Roman"/><w:sz w:val="28"/><w:szCs w:val="28"/>'
        '</w:rPr><w:t>cells</w:t></w:r></w:p>',
        head + '<w:r><w:rPr><w:b/><w:bCs/></w:rPr><w:t>Bold</w:t></w:r><w:r><w:rPr><w:b/></w:rPr>'
        '<w:t xml:space="preserve"> words</w:t></w:r><w:r><w:rPr><w:iCs/></w:rPr>'
        '<w:t xml:space="preserve"> and plain text here</w:t></w:r></w:p>',
        head + '<w:r><w:t xml:space="preserve">Read the </w:t></w:r><w:hyperlink r:id="rIdLINK1"><w:r><w:rPr><w:b/>'
        '</w:rPr><w:t>WHO</w:t></w:r><w:r><w:t xml:space="preserve"> classification site</w:t></w:r></w:hyperlink>'
        '<w:r><w:t xml:space="preserve"> now and then again</w:t></w:r></w:p>',
        head + '<w:r><w:lastRenderedPageBreak/><w:t>Second page text</w:t></w:r></w:p>',
        head + '<w:r><w:t>Text kept</w:t></w:r><w:ins w:id="1" w:author="a" w:date="2024-01-01T00:00:00Z"><w:r>'
        '<w:t>new</w:t></w:r></w:ins></w:p>',
        head + '<w:r><w:t>Text kept</w:t></w:r><w:del w:id="2" w:author="a" w:date="2024-01-01T00:00:00Z"><w:r>'
        '<w:delText>old</w:delText></w:r></w:del></w:p>',
        head + '<w:pPr><w:rPr><w:b/><w14:ligatures w14:val="all"/></w:rPr></w:pPr><w:r><w:t>Text mark</w:t></w:r>'
        '</w:p>',
        head + '<w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:t>Centred text</w:t></w:r></w:p>',
        head + '<w:r><w:t xml:space="preserve">Intro text </w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> TOC </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/>'
        '</w:r><w:r><w:t>Result text</w:t></w:r></w:p>',
        head + '<w:r><w:t xml:space="preserve">Date: </w:t></w:r><w:sdt><w:sdtPr/><w:sdtContent><w:r>'
        '<w:t>Click here to enter a date</w:t></w:r></w:sdtContent></w:sdt></w:p>',
        head + '<w:r><w:t xml:space="preserve">See </w:t></w:r><w:fldSimple w:instr=" REF _Ref1 "><w:r>'
        '<w:t>Table 1</w:t></w:r></w:fldSimple><w:r><w:t xml:space="preserve"> for details</w:t></w:r></w:p>',
        head + '<w:hyperlink w:anchor="_Toc1" w:history="1"><w:r><w:t>Introduction</w:t></w:r><w:r><w:tab/></w:r>'
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> PAGEREF _Toc1 '
        '</w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>3</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:hyperlink></w:p>',
        head + '<w:r><w:t xml:space="preserve">Visit </w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> HYPERLINK "https://www.who.int" </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>the WHO website</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r><w:r><w:t xml:space="preserve"> today</w:t></w:r></w:p>',
        head + '<w:r><w:t xml:space="preserve">Grades of the tumour: </w:t></w:r><w:r><w:rPr><w:b/></w:rPr>'
        '<w:t>a</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t xml:space="preserve"> </w:t></w:r><w:r><w:rPr><w:b/></w:rPr>'
        '<w:t>b</w:t></w:r></w:p>',
        head + '<w:r><w:t>A</w:t></w:r><w:r><w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol"/></w:rPr>'
        '<w:t>\uf062</w:t></w:r><w:r><w:t xml:space="preserve"> amyloid plaque</w:t></w:r></w:p>',
        head + '<w:r><w:t xml:space="preserve">Comment </w:t></w:r><!-- a comment --><w:r><w:t>here</w:t>'
        '<!-- in run --></w:r><w:customXml w:element="x"><w:customXmlPr/><w:r><w:rPr><w:i/></w:rPr>'
        '<w:t>custom</w:t></w:r></w:customXml><w:r><w:t xml:space="preserve"> tail \u00a0 text\u2003</w:t></w:r>'
        '<w:r><w:rPr><w:u w:val="single"/></w:rPr><w:t>\ufeff</w:t></w:r><w:r><w:t>closing words</w:t></w:r></w:p>',
        head + '<w:r><w:rPr><w:b/></w:rPr><w:t>Bold</w:t></w:r><w:r><w:rPr><w:u w:val="single"/></w:rPr>'
        '<w:t xml:space="preserve">\u2028</w:t></w:r><w:r><w:t xml:space="preserve">\u0085</w:t></w:r>'
        '<w:r><w:rPr><w:b/></w:rPr><w:t>again</w:t></w:r><w:r><w:t xml:space="preserve"> and plain words</w:t></w:r>'
        '</w:p>',
        head + '<w:r><w:t xml:space="preserve">Line one</w:t></w:r><w:r><w:br/></w:r><w:r><w:rPr><w:b/>'
        '</w:rPr><w:t>Bold</w:t><w:tab/><w:t>after tab</w:t></w:r><w:r><w:t xml:space="preserve"> and the '
        'end of a long plain line</w:t></w:r></w:p>',
        head + '\n  <w:r>\n    <w:rPr>\n      <w:b/>\n    </w:rPr>\n    <w:t>Pretty</w:t>\n    <w:tab/>\n'
        '    <w:t xml:space="preserve"> printed </w:t>\n  </w:r>\n  <w:hyperlink r:id="rIdLINK1">\n    <w:r>\n'
        '      <w:t>linked words</w:t>\n    </w:r>\n  </w:hyperlink>\n  <w:r>\n    <w:t xml:space="preserve"> and the '
        'plain end</w:t>\n  </w:r>\n</w:p>',
        head + '<w:r><w:rPr><w:b/></w:rPr><w:t>Alpha</w:t></w:r><w:r><w:rPr><w:i/></w:rPr><w:t>Gamma</w:t></w:r></w:p>',
    ]


def _codec_a_paragraphs():
    """The DrawingML paragraphs of helper/tests/test_codec.py (synthetic text)."""
    head = '<a:p xmlns:a="%s">' % A

    def run(text, attrs=""):
        return '<a:r><a:rPr lang="en-US"%s/><a:t>%s</a:t></a:r>' % (attrs, text)
    return [
        head + '\n  <a:pPr algn="l"/>\n  <a:r><a:rPr lang="en-US" sz="1800" dirty="0"/><a:t>In </a:t></a:r>\n'
        '  <a:r><a:rPr lang="en-US" sz="1800" b="1" dirty="0"/><a:t>NF2-altered</a:t></a:r>\n'
        '  <a:r><a:rPr lang="en-US" sz="1800" dirty="0"/><a:t> meningiomas, </a:t></a:r>\n'
        '  <a:r><a:rPr lang="en-US" sz="2000"/><a:t>chromosome</a:t></a:r>\n  <a:br><a:rPr lang="en-US"/></a:br>\n'
        '  <a:r><a:rPr lang="en-US" sz="1800" i="1"/><a:t>1p loss</a:t></a:r>\n'
        '  <a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum"><a:rPr lang="en-US"/><a:t>3</a:t>'
        '</a:fld>\n  <a:endParaRPr lang="en-US" sz="1800"/>\n</a:p>',
        head + run("Smith A. et al. ") + run("J Clin", ' i="1"') + run(" ", ' i="1" err="1"') + run("Pathol", ' i="1"')
        + '</a:p>',
        head + '<a:r><a:rPr lang="en-US"><a:latin typeface="Symbol" pitchFamily="18" charset="2"/></a:rPr>'
        '<a:t>\uf061</a:t></a:r><a:r><a:rPr lang="en-US"><a:latin typeface="Arial"/></a:rPr>'
        '<a:t>-synuclein immunoreactivity</a:t></a:r></a:p>',
        head + '<a:r><a:rPr lang="en-US"><a:sym typeface="Wingdings"/></a:rPr><a:t>\uf0e8 </a:t></a:r>'
        '<a:r><a:rPr lang="en-US"><a:sym typeface="Wingdings"/></a:rPr><a:t>Lateral neck LNs</a:t></a:r>'
        '<a:r><a:rPr lang="en-US"/><a:t> (14.6%)</a:t></a:r></a:p>',
        head + run("Hello") + run(" ", ' u="sng"') + run("world") + '</a:p>',
        head + run("Results of the study: ") + run("a", ' u="sng"') + run(" ", ' b="1"') + run("b", ' u="sng"')
        + '</a:p>',
        head + '<a:r><a:rPr lang="en-US" sz="1800"><a:latin typeface="Arial"/></a:rPr><a:t>Tumour </a:t></a:r>'
        '<a:r><a:rPr lang="en-US" sz="2400"><a:latin typeface="Calibri"/></a:rPr><a:t>cells</a:t></a:r></a:p>',
        head + run("Spindle cell ") + run("lesion", ' spc="-10"') + run(" with ") + run("myxoid", ' spc="-20"')
        + run(" stroma ") + run("(SCL-", ' spc="-35"') + run("MS)", ' spc="-20"') + '</a:p>',
        head + run("Loss of ", ' spc="-10"') + run("ABC1", ' i="1" spc="-45"') + run("-mutated", ' spc="-45"')
        + run(" cells") + '</a:p>',
        head + run("Updated ") + '<a:fld id="{1A2B3C4D-0000-4000-8000-000000000001}" type="datetime8">'
        '<a:rPr lang="en-US"/><a:t>Tuesday, October 13</a:t></a:fld></a:p>',
        head + run("Text with a math sign ") + run("\U0001d6fc", ' i="1"') + run(" and an emoji \U0001f600 end")
        + '</a:p>',
        head + '<a:endParaRPr lang="en-US"/></a:p>',
        head + '\n    <a:r>\n      <a:rPr lang="en-US" b="1"/>\n      <a:t>Pretty printed</a:t>\n    </a:r>\n'
        '    <a:r>\n      <a:rPr lang="en-US"/>\n      <a:t> runs with tails</a:t>\n    </a:r>\n    <a:br/>\n'
        '    <a:r>\n      <a:rPr lang="en-US" i="1"/>\n      <a:t>end</a:t>\n    </a:r>\n  </a:p>',
        head + run("Alpha", ' b="1"') + run("Gamma", ' i="1"') + '</a:p>',
    ]


def _insert_w_paragraphs(d, paras, links):
    body = d.element.body
    sect = body.find("{%s}sectPr" % W)
    for xml in paras:
        for key, rid in links.items():
            xml = xml.replace(key, rid)
        el = etree.fromstring(xml)
        if sect is not None:
            sect.addprevious(el)
        else:
            body.append(el)


def make_codec_docx(path):
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE
    d = docx.Document()
    d.add_paragraph("Codec cases follow")
    rid = d.part.relate_to("https://www.who.int", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    _insert_w_paragraphs(d, _codec_w_paragraphs(), {"rIdLINK1": rid})
    d.save(path)
    return path


def make_codec_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(6))
    tx = box.text_frame._txBody
    for p in tx.findall("{%s}p" % A):
        tx.remove(p)
    for xml in _codec_a_paragraphs():
        tx.append(etree.fromstring(xml))
    prs.save(path)
    return path


_SP_BODIES = [
    '<a:bodyPr xmlns:a="%s"><a:prstTxWarp prst="textNoShape"><a:avLst/></a:prstTxWarp>'
    '<a:normAutofit fontScale="90000"/><a:extLst/></a:bodyPr>',
    '<a:bodyPr xmlns:a="%s" wrap="none"><a:spAutoFit/></a:bodyPr>',
    '<a:bodyPr xmlns:a="%s"><a:normAutofit fontScale="62.5%%"/></a:bodyPr>',
    '<a:bodyPr xmlns:a="%s"><a:normAutofit fontScale="90000" lnSpcReduction="20000"/></a:bodyPr>',
    '<a:bodyPr xmlns:a="%s" wrap="square">\n  <a:noAutofit/>\n</a:bodyPr>',
    '<a:bodyPr xmlns:a="%s"/>',
    '<a:bodyPr xmlns:a="%s" wrap="square">\n  <a:prstTxWarp prst="textNoShape">\n    <a:avLst/>\n  </a:prstTxWarp>\n'
    '  <a:normAutofit fontScale="85000"/>\n  <a:scene3d><a:camera prst="orthographicFront"/>'
    '<a:lightRig rig="threePt" dir="t"/></a:scene3d>\n</a:bodyPr>',
]


def make_fit_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    texts = ["Short English line", "Label", "Scaled text box", "Spacing kept here", "Plain box text",
             "Grows a lot", "Warped text"]
    for n, (body, text) in enumerate(zip(_SP_BODIES, texts)):
        box = s.shapes.add_textbox(Inches(0.5), Inches(0.5 + n), Inches(4), Inches(0.6))
        box.text_frame.text = text
        tx = box.text_frame._txBody
        old = tx.find("{%s}bodyPr" % A)
        new = etree.fromstring(body % A)
        old.addprevious(new)
        tx.remove(old)
    second = s.shapes.add_textbox(Inches(5), Inches(0.5), Inches(4), Inches(1)).text_frame
    second.text = "First line of a box"
    second.add_paragraph().text = "Second line of the same box"
    table = s.shapes.add_table(2, 2, Inches(5), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Marker"
    table.cell(0, 1).text = "Result of the stain"
    table.cell(1, 0).text = "Ki-67"
    table.cell(1, 1).text = "Low"
    # one long word: "VI_" adds 3 of 45 characters, a length ratio of 1.07 (just above the 1.05 threshold)
    near = s.shapes.add_textbox(Inches(5), Inches(5), Inches(4), Inches(0.6)).text_frame
    near.text = "Pneumonoultramicroscopicsilicovolcanoconiosis"
    prs.save(path)
    return path


def make_chart_pptx(path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Cases by grade"
    data = CategoryChartData()
    data.categories = ["Grade 1", "Grade 2", "Grade 3"]
    data.add_series("Cases", (12, 7, 3))
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(6), Inches(4),
                               data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Number of cases per grade"
    s.notes_slide.notes_text_frame.text = "Explain the chart briefly"
    prs.save(path)
    return path


def make_comments_docx(path):
    import docx
    d = docx.Document()
    p = d.add_paragraph("Body text with a remark")
    d.add_comment(p.runs, text="Check this staining result", author="Reviewer", initials="R")
    d.add_paragraph("Second paragraph of the body")
    d.save(path)
    return path


def _deck3(path):
    import pptx
    prs = pptx.Presentation()
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Title number %d" % (i + 1)
    prs.slides[2].notes_slide.notes_text_frame.text = "Note on the third slide"
    prs.save(path)
    return path


def _orphan(path):
    _deck3(path)
    pkg = OoxmlPackage(path)
    root = pkg.root("ppt/presentation.xml")
    lst = root.find(".//{%s}sldIdLst" % P)
    for i, el in enumerate(list(lst)):
        if i == 2:
            lst.remove(el)
    pkg.mark_dirty("ppt/presentation.xml")
    tmp = Path(str(path) + ".tmp.pptx")
    pkg.save(tmp)
    tmp.replace(path)
    return path


def _notes_order(path):
    import pptx
    prs = pptx.Presentation()
    for i in range(3):
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Slide title %d" % (i + 1)
    prs.slides[2].notes_slide.notes_text_frame.text = "Notes for the third slide"
    prs.slides[0].notes_slide.notes_text_frame.text = "Notes for the first slide"
    prs.save(path)
    return path


def _patch_zip(src, dst, edits):
    """Copy package src to dst with each part of `edits` rewritten as text by its function (raw XML: declarations
    stay as written)."""
    with zipfile.ZipFile(src) as z:
        infos = z.infolist()
        data = {i.filename: z.read(i.filename) for i in infos}
    for part, edit in edits.items():
        text = data[part].decode("utf-8")
        new = edit(text)
        if new == text:
            raise ValueError("edit changed nothing in %s" % part)
        data[part] = new.encode("utf-8")
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for i in infos:
            z.writestr(i, data[i.filename])
    return dst


def _add_rel(rel):
    return lambda t: t.replace("</Relationships>", rel + "</Relationships>", 1)


def _notes_shared(path):
    """deck3 whose slide 1 also relates to the notes slide of slide 3 (the helper reads targets, not relationship
    types: the notes are labelled by the first slide that points at them). An unknown relationship type keeps the
    deck valid for PowerPoint, which refuses one notes slide shared by two slides."""
    _deck3(path)
    rel = ('<Relationship Id="rId99" Type="http://example.invalid/vitranslate-test/related-notes" '
           'Target="../notesSlides/notesSlide1.xml"/>')
    tmp = Path(str(path) + ".tmp.pptx")
    _patch_zip(path, tmp, {"ppt/slides/_rels/slide1.xml.rels": _add_rel(rel)})
    tmp.replace(path)
    return path


# Paragraphs that carry their own (redundant) namespace declarations, inserted as raw XML so the declarations stay:
# lxml drops a declaration the new parent already has whenever it moves or copies an element.
_ADV_DOCX = {
    "sdt": '<w:p><w:r><w:t xml:space="preserve">Intro words </w:t></w:r><w:sdt xmlns:w="%s"><w:sdtContent><w:r>'
           '<w:t>locked</w:t></w:r></w:sdtContent></w:sdt><w:r><w:t xml:space="preserve"> and the end text</w:t></w:r>'
           '</w:p>' % W,
    "anchor": '<w:p><w:bookmarkStart xmlns:w="%s" w:id="5" w:name="x"/><w:r><w:t>Anchored words here</w:t></w:r>'
              '<w:bookmarkEnd w:id="5"/></w:p>' % W,
    "rpr": '<w:p><w:r><w:rPr xmlns:w="%s"><w:b/></w:rPr><w:t>Bold</w:t></w:r><w:r>'
           '<w:t xml:space="preserve"> and plain words</w:t></w:r></w:p>' % W,
    "hyperlink": '<w:p><w:r><w:t xml:space="preserve">Read the </w:t></w:r><w:hyperlink xmlns:r="%s" r:id="rIdAdv1">'
                 '<w:r><w:t>linked words</w:t></w:r></w:hyperlink><w:r><w:t xml:space="preserve"> now</w:t></w:r>'
                 '</w:p>' % R,
    "tab_in_run": '<w:p><w:r><w:t>Before</w:t><w:tab xmlns:w="%s"/><w:t>after the tab</w:t></w:r></w:p>' % W,
}
_ADV_PPTX = {
    "br": '<a:p><a:r><a:rPr lang="en-US"/><a:t>Start words</a:t></a:r><a:br xmlns:a="%s"/><a:r><a:rPr lang="en-US"/>'
          '<a:t>more words</a:t></a:r></a:p>' % A,
    "rpr": '<a:p><a:r><a:rPr xmlns:a="%s" lang="en-US" b="1"/><a:t>Bold words</a:t></a:r><a:r><a:rPr lang="en-US"/>'
           '<a:t> plain tail text</a:t></a:r></a:p>' % A,
}


_ADV_LINK_REL = ('<Relationship Id="rIdAdv1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                 'relationships/hyperlink" Target="https://www.who.int" TargetMode="External"/>')


def _adv_docx(para):
    edits = {"word/document.xml": lambda t: t.replace("<w:sectPr", para + "<w:sectPr", 1)}
    if "rIdAdv1" in para:
        edits["word/_rels/document.xml.rels"] = _add_rel(_ADV_LINK_REL)
    return lambda path: _patch_zip(FIXTURES / "refs_column.docx", path, edits)


def _adv_pptx(para):
    return lambda path: _patch_zip(FIXTURES / "codec_cases.pptx", path,
                                   {"ppt/slides/slide1.xml": lambda t: t.replace("</p:txBody>", para + "</p:txBody>", 1)})


class _Unseekable:
    """A write-only stream: zipfile then writes data descriptors (flag bit 3)."""

    def __init__(self):
        self.buf = io.BytesIO()

    def write(self, b):
        return self.buf.write(b)

    def flush(self):
        pass


def _zip_entries(src):
    with zipfile.ZipFile(src) as z:
        return [(i, z.read(i.filename)) for i in z.infolist()]


def _zip64(path):
    """doc.docx with zip64 extra fields on every entry and a zip64 end of central directory record."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in _zip_entries(FIXTURES / "doc.docx"):
            with z.open(zipfile.ZipInfo(info.filename, info.date_time), "w", force_zip64=True) as f:
                f.write(data)
    return path


def _datadesc(path):
    """doc.docx written by a streaming writer: data descriptors after every entry."""
    s = _Unseekable()
    with zipfile.ZipFile(s, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in _zip_entries(FIXTURES / "doc.docx"):
            z.writestr(zipfile.ZipInfo(info.filename, info.date_time), data, compress_type=zipfile.ZIP_DEFLATED)
    Path(path).write_bytes(s.buf.getvalue())
    return path


def _comment_extra(path):
    """doc.docx with an archive comment and an extended-timestamp extra field on every entry."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.comment = b"synthetic archive comment"
        for info, data in _zip_entries(FIXTURES / "doc.docx"):
            zi = zipfile.ZipInfo(info.filename, info.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.extra = struct.pack("<HHBI", 0x5455, 5, 1, 1700000000)
            z.writestr(zi, data)
    return path


def _docx_build(path, build):
    import docx
    d = docx.Document()
    build(d)
    d.save(path)
    return path


def _para(d, text, how):
    from fixtures import set_outline_level
    kind, value = how
    if kind == "style":
        return d.add_paragraph(text, style=value)
    if kind == "style+outline":
        return set_outline_level(d.add_paragraph(text, style=value[0]), value[1])
    p = d.add_paragraph(text)
    if kind == "outline":
        set_outline_level(p, value)
    else:
        p._p.style = value
    return p


def synthetic_builders():
    from fixtures import add_text_box_paragraph, make_docx, make_pptx, reorder_slides, set_outline_level

    def refs_column(d):
        lab = d.add_table(rows=2, cols=3)
        for c, text in enumerate(("Test", "Result", "Reference")):
            lab.cell(0, c).text = text
        for c, text in enumerate(("Ki-67 index", "12 percent", "Below 4 percent")):
            lab.cell(1, c).text = text
        d.add_paragraph("Final diagnosis: atypical meningioma.")

    def refs_textbox(d):
        add_text_box_paragraph(d, "References")
        d.add_paragraph("Final diagnosis: atypical meningioma.")

    def refs_region(style, outline):
        def build(d):
            d.add_paragraph("Body text before the list")
            head = d.add_paragraph("References", style=style)
            if outline is not None:
                set_outline_level(head, outline)
            d.add_paragraph("Louis DN et al. Neuro Oncol 2021")
            d.add_heading("Appendix A", level=1)
            d.add_paragraph("Staining protocol for the lab")
        return build

    h1, h2, h3 = ("style", "Heading 1"), ("style", "Heading 2"), ("style", "Heading 3")

    def refs_sub(stop, sub, appendix):
        def build(d):
            _para(d, "References", stop)
            _para(d, "Primary sources", sub)
            d.add_paragraph("Louis DN et al. Neuro Oncol 2021")
            _para(d, "Reviews", sub)
            d.add_paragraph("Perry A et al. Brain Pathol 2020")
            _para(d, "Appendix A", appendix)
            d.add_paragraph("Staining protocol for the lab")
        return build

    def refs_nolevel(d):
        _para(d, "References", ("id", "ReferencesHeading"))
        d.add_paragraph("Louis DN et al. Neuro Oncol 2021")
        _para(d, "Appendix A", h3)
        d.add_paragraph("Staining protocol for the lab")

    # a back-matter title closes the references region whatever its style or level (his report 2026-10-15)
    def refs_backmatter(stop, title):
        def build(d):
            _para(d, "References", stop)
            d.add_paragraph("Louis DN et al. Neuro Oncol 2021")
            d.add_paragraph("Funding. J Synth Path. 2020;1:1-9.")
            _para(d, "Acknowledgements", title)
            d.add_paragraph("Staining protocol for the lab")
        return build

    def refs_plain_inside(d):
        d.add_paragraph("References")
        d.add_paragraph("Louis DN et al. Neuro Oncol 2021")
        d.add_paragraph("Perry A, Example B. A synthetic study of the lab. Brain Pathol 2020;30:1-9.")
        d.add_paragraph("Staining protocol for the lab")

    out = {
        "deck.pptx": make_pptx,
        "deck_reordered.pptx": lambda p: reorder_slides(make_pptx(p), [4, 3, 2, 1, 0]),
        "doc.docx": make_docx,
        "deck3.pptx": _deck3,
        "deck3_orphan.pptx": _orphan,
        "notes_order.pptx": _notes_order,
        "codec_cases.docx": make_codec_docx,
        "codec_cases.pptx": make_codec_pptx,
        "fit_cases.pptx": make_fit_pptx,
        "chart.pptx": make_chart_pptx,
        "comments.docx": make_comments_docx,
        "refs_column.docx": lambda p: _docx_build(p, refs_column),
        "refs_textbox.docx": lambda p: _docx_build(p, refs_textbox),
        "refs_nolevel.docx": lambda p: _docx_build(p, refs_nolevel),
        "refs_backmatter_plain.docx": lambda p: _docx_build(p, refs_backmatter(("style", "Normal"), ("style", "Normal"))),
        "refs_backmatter_deeper.docx": lambda p: _docx_build(p, refs_backmatter(h1, h3)),
        "refs_plain_inside.docx": lambda p: _docx_build(p, refs_plain_inside),
        "notes_shared.pptx": _notes_shared,
        "zip64.docx": _zip64,
        "zip_datadesc.docx": _datadesc,
        "zip_comment_extra.docx": _comment_extra,
    }
    for key, para in _ADV_DOCX.items():
        out["adv_%s.docx" % key] = _adv_docx(para)
    for key, para in _ADV_PPTX.items():
        out["adv_%s.pptx" % key] = _adv_pptx(para)
    for k, (style, outline) in enumerate([("Heading 1", None), ("Title", None), (None, 0), (None, None), (None, 9)]):
        out["refs_region_%d.docx" % (k + 1)] = (lambda s, o: lambda p: _docx_build(p, refs_region(s, o)))(style,
                                                                                                         outline)
    # case 7: outline level 8 is the deepest heading level (9 = body text)
    for k, case in enumerate([(h1, h2, h1), (("outline", 0), ("outline", 1), ("outline", 0)),
                              (("style", "Title"), h2, h1), (("outline", 0), h2, h1),
                              (h1, h2, ("id", "AppendixHeading")), (("style+outline", ("Heading 2", 0)), h2, h1),
                              (("outline", 8), ("style", "Normal"), ("outline", 8))]):
        out["refs_sub_%d.docx" % (k + 1)] = (lambda c: lambda p: _docx_build(p, refs_sub(*c)))(case)
    return out


OPTIONS = {"comments.docx": {"comments": True}}


def build_fixtures(folder):
    folder.mkdir(parents=True, exist_ok=True)
    names = []
    for name, build in synthetic_builders().items():
        path = folder / name
        build(path)
        names.append(name)
    return names


def main(argv):
    if not argv:
        print(__doc__)
        return 2
    cmd = argv[0]
    if cmd == "fixtures":
        names = build_fixtures(FIXTURES)
        for name in names:
            data = dump(FIXTURES / name, **OPTIONS.get(name, {}))
            write_json(data, FIXTURES / (name + ".json"))
            print("%-24s items %4d  segments %4d" % (name, len(data["items"]),
                                                    sum(1 for i in data["items"] if i["seg"])))
        write_json(sorted(name + ".json" for name in names), FIXTURES / "index.json")  # for the browser check page
        write_json(lxml_cases(), FIXTURES.parent / "lxml_cases.json")
        return 0
    if cmd == "dump" and len(argv) >= 3:
        out = Path(argv[1])
        out.mkdir(parents=True, exist_ok=True)
        for f in argv[2:]:
            write_json(dump(f), out / (Path(f).name + ".json"))
            print("dumped", Path(f).name.encode("ascii", "replace").decode("ascii"))
        return 0
    if cmd == "corpus" and len(argv) == 2:
        out = Path(argv[1]).resolve()
        if ROOT.resolve() in out.parents or out == ROOT.resolve():
            print("refused: corpus dumps hold private text and must stay outside the repository")
            return 2
        out.mkdir(parents=True, exist_ok=True)
        listing = ROOT / "private" / "corpus.txt"
        paths = [Path(s.strip()) for s in listing.read_text(encoding="utf-8").splitlines() if s.strip()]
        n = 0
        for path in paths:
            if path.suffix.lower() not in (".docx", ".pptx"):
                continue
            n += 1
            data = dump(path, source=str(path))
            write_json(data, out / ("corpus_%d.json" % n))
            print("corpus_%d %s items %d segments %d" % (n, data["kind"], len(data["items"]),
                                                          sum(1 for i in data["items"] if i["seg"])))
        return 0
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
